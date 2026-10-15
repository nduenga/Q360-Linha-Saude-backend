"""Recalcula o separador `Calculador` do modelo actuarial com outro segmento.

O Excel só guarda os valores calculados do cenário em que foi gravado (SAU).
Para obter os valores esperados de outro cenário sem abrir o Excel, este
script grava uma cópia com `Calculador!C5` alterado e avalia as fórmulas do
próprio livro com o pycel. Ferramenta de desenvolvimento: não é uma
dependência do backend. Uso (num ambiente à parte):

    python -m venv /tmp/xl && /tmp/xl/bin/pip install pycel "openpyxl<3.1"
    /tmp/xl/bin/python tools/recalc_calculador.py SAG > tests/quote/fixtures/calculador_sag_expected.json

Validação: com SAU o resultado tem de coincidir com os valores guardados
pelo Excel (o script falha se a diferença exceder 1e-6).
"""
from __future__ import annotations

import hashlib
import json
import math
import sys
import tempfile
import types
from pathlib import Path

from openpyxl import load_workbook
from pycel import ExcelCompiler

WORKBOOK = Path(__file__).resolve().parents[1] / "apps/business_config/seed/Modelo_Actuarial_Saude_SAU_SAG_Angola.xlsx"
CELLS = ["C7", "K25", "K26", "K27", "K28", "C59", "C60", "C61", "C62", "C63", "C64", "C65", "C66", "C70", "C71",
         "C75", "C76", "C77", "C78", "C79", "C84", "C85", "C86", "C87", "E91", "E92", "E93", "E94"]


def normsdist(z):
    if isinstance(z, str):  # o pycel avalia os dois ramos do IF: um erro (#NUM!) propaga-se, como no Excel
        return z
    return (1 + math.erf(float(z) / math.sqrt(2))) / 2


def evaluate(segment: str) -> dict:
    plugin = types.ModuleType("q360_xl_plugin")
    plugin.normsdist = normsdist
    sys.modules["q360_xl_plugin"] = plugin
    with tempfile.TemporaryDirectory() as tmp:
        copy = Path(tmp) / "calculador.xlsx"
        workbook = load_workbook(WORKBOOK)
        workbook["Calculador"]["C5"] = segment
        workbook.save(copy)
        compiler = ExcelCompiler(filename=str(copy), plugins=["q360_xl_plugin"])
        return {cell: compiler.evaluate(f"Calculador!{cell}") for cell in CELLS}


def main() -> None:
    segment = sys.argv[1] if len(sys.argv) > 1 else "SAG"
    cached = load_workbook(WORKBOOK, data_only=True)["Calculador"]
    check = evaluate("SAU")
    worst = max(abs(float(v) - float(cached[k].value)) for k, v in check.items() if k != "C7")
    if worst > 1e-6:
        raise SystemExit(f"O recálculo SAU difere do Excel em {worst}: método não fiável.")
    print(json.dumps({
        "_provenance": {"workbookSha256": hashlib.sha256(WORKBOOK.read_bytes()).hexdigest(),
                        "change": f"Calculador!C5 = {segment}", "sauCheckMaxDiff": worst},
        "cells": evaluate(segment),
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
