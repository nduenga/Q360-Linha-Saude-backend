"""`/api/v1/platform/system-parameters` — camada de configuração geral do
fornecedor: parâmetros que valem para todos os tenants (estruturais,
impostos legais e valores por omissão/limites dos operacionais). Cada
alteração é uma versão proposta por uma conta da plataforma e aprovada por
outra; a versão em vigor é a publicada mais recente cuja vigência começou."""
from __future__ import annotations

from io import BytesIO

from django.http import HttpResponse
from rest_framework.parsers import MultiPartParser

from rest_framework import serializers, status
from rest_framework.response import Response
from rest_framework.views import APIView

from apps.foundation.application import platform_parameter_service as service
from apps.foundation.application import platform_reference_service as reference_service
from apps.foundation.interfaces.middleware import get_correlation_id, requires_idempotency
from apps.platform_admin.application.audit_service import record_platform_audit
from apps.platform_admin.interfaces.permissions import IsPlatformAdmin


class PlatformParameterProposalSerializer(serializers.Serializer):
    value = serializers.JSONField()
    minimum = serializers.JSONField(required=False, allow_null=True, default=None)
    maximum = serializers.JSONField(required=False, allow_null=True, default=None)
    effectiveFrom = serializers.DateTimeField(required=False, allow_null=True, default=None)
    reason = serializers.CharField(required=False, allow_null=True, allow_blank=True, default=None)
    legalReference = serializers.CharField(required=False, allow_null=True, allow_blank=True, default=None, max_length=1000)
    sourceUrl = serializers.CharField(required=False, allow_null=True, allow_blank=True, default=None, max_length=500)


class PlatformParameterDecisionSerializer(serializers.Serializer):
    note = serializers.CharField(required=False, allow_null=True, allow_blank=True, default=None)


class PlatformParameterRollbackSerializer(serializers.Serializer):
    versionNo = serializers.IntegerField(min_value=1)
    reason = serializers.CharField(required=False, allow_null=True, allow_blank=True, default=None)


def _actor_id(request) -> str:
    return request.user.actor.actor_id


def _audit(request, action: str, code: str, payload: dict) -> None:
    record_platform_audit(
        actor_id=_actor_id(request), action=action, entity_type="PlatformParameter", entity_id=code,
        correlation_id=get_correlation_id(),
        metadata={"versionNo": payload.get("versionNo"), "value": payload.get("value"),
                  "minimum": payload.get("minimum"), "maximum": payload.get("maximum"),
                  "effectiveFrom": payload.get("effectiveFrom"), "reason": payload.get("reason"),
                  "legalReference": payload.get("legalReference"),
                  "decisionNote": payload.get("decisionNote")},
    )


class PlatformParameterListView(APIView):
    """Every parameter as the platform sees it: level, value in effect,
    code default, limits imposed on tenants and pending versions.
    Filters: `?level=PLATFORM_MANDATORY|PLATFORM_DEFAULT`, `?category=`."""

    permission_classes = [IsPlatformAdmin]

    def get(self, request):
        return Response(service.list_parameters(
            level=request.query_params.get("level"), category=request.query_params.get("category"),
        ))


class PlatformParameterDetailView(APIView):
    permission_classes = [IsPlatformAdmin]

    def get(self, request, code):
        return Response(service.describe(code))


@requires_idempotency
class PlatformParameterVersionsView(APIView):
    """GET: version history. POST: propose a new version (pending approval
    by another platform account)."""

    permission_classes = [IsPlatformAdmin]

    def get(self, request, code):
        return Response(service.list_versions(code))

    def post(self, request, code):
        body = PlatformParameterProposalSerializer(data=request.data)
        body.is_valid(raise_exception=True)
        data = body.validated_data
        created = service.propose(
            code=code, value=data["value"], minimum=data["minimum"], maximum=data["maximum"],
            effective_from=data["effectiveFrom"], reason=data["reason"], actor_id=_actor_id(request),
            legal_reference=data["legalReference"], source_url=data["sourceUrl"],
        )
        _audit(request, "platform_parameter.proposed", code, created)
        return Response(created, status=status.HTTP_201_CREATED)


@requires_idempotency
class PlatformParameterApproveView(APIView):
    permission_classes = [IsPlatformAdmin]

    def post(self, request, code, version_no):
        body = PlatformParameterDecisionSerializer(data=request.data)
        body.is_valid(raise_exception=True)
        approved = service.approve(code=code, version_no=version_no, actor_id=_actor_id(request),
                                   note=body.validated_data["note"])
        _audit(request, "platform_parameter.approved", code, approved)
        return Response(approved)


@requires_idempotency
class PlatformParameterRejectView(APIView):
    permission_classes = [IsPlatformAdmin]

    def post(self, request, code, version_no):
        body = PlatformParameterDecisionSerializer(data=request.data)
        body.is_valid(raise_exception=True)
        rejected = service.reject(code=code, version_no=version_no, actor_id=_actor_id(request),
                                  note=body.validated_data["note"])
        _audit(request, "platform_parameter.rejected", code, rejected)
        return Response(rejected)


@requires_idempotency
class PlatformParameterRollbackView(APIView):
    permission_classes = [IsPlatformAdmin]

    def post(self, request, code):
        body = PlatformParameterRollbackSerializer(data=request.data)
        body.is_valid(raise_exception=True)
        created = service.rollback(code=code, version_no=body.validated_data["versionNo"],
                                   reason=body.validated_data["reason"], actor_id=_actor_id(request))
        _audit(request, "platform_parameter.proposed", code, created)
        return Response(created, status=status.HTTP_201_CREATED)


class PlatformFormulaEngineView(APIView):
    """O motor de fórmulas e regras visto pelo fornecedor (sem tenant):
    funções, operadores, limites e variáveis do sistema de cada contexto, e
    os campos/ações das regras «Quando… então…». Só leitura: mudar o motor é
    uma nova versão do sistema; as fórmulas e regras de negócio são de cada
    tenant (`/business-config/formulas`)."""

    permission_classes = [IsPlatformAdmin]

    def get(self, request):
        from apps.business_config.domain import custom_rules
        from apps.business_config.domain.formula_engine import MAX_EXPONENT, MAX_EXPRESSION_LENGTH, MAX_NODES
        from apps.business_config.domain.formulas import formula_catalog
        from apps.business_config.domain.product_setup import SYSTEM_VARIABLES as PRODUCT_SYSTEM_VARIABLES
        from apps.foundation.domain.system_parameters import DEFINITIONS, PLATFORM_CONSTANTS

        return Response({
            "pricingFormulas": formula_catalog(),
            "productFormulas": {"systemVariables": [{"name": k, "label": v} for k, v in PRODUCT_SYSTEM_VARIABLES.items()]},
            "rules": custom_rules.field_catalog(),
            "platformConstants": [
                {"name": name, "parameter": code, "unit": DEFINITIONS[code].unit,
                 "description": DEFINITIONS[code].description}
                for name, code in sorted(PLATFORM_CONSTANTS.items())
            ],
            "limits": {"maxExpressionLength": MAX_EXPRESSION_LENGTH, "maxNodes": MAX_NODES, "maxExponent": MAX_EXPONENT},
        })


# ------------------------------------------------------- tabelas de referência
class ReferenceTableProposalSerializer(serializers.Serializer):
    code = serializers.CharField(max_length=60)
    name = serializers.CharField(max_length=200)
    description = serializers.CharField(required=False, allow_null=True, allow_blank=True, default=None)
    keyColumn = serializers.CharField(max_length=60)
    columns = serializers.ListField(child=serializers.DictField())
    rows = serializers.ListField(child=serializers.DictField())
    source = serializers.CharField(required=False, allow_null=True, allow_blank=True, default=None, max_length=1000)
    effectiveFrom = serializers.DateTimeField(required=False, allow_null=True, default=None)
    reason = serializers.CharField(required=False, allow_null=True, allow_blank=True, default=None)


def _table_audit(request, action: str, payload: dict) -> None:
    record_platform_audit(
        actor_id=_actor_id(request), action=action, entity_type="PlatformReferenceTable", entity_id=payload.get("code"),
        correlation_id=get_correlation_id(),
        metadata={"versionNo": payload.get("versionNo"), "rowCount": payload.get("rowCount"),
                  "source": payload.get("source"), "reason": payload.get("reason"),
                  "decisionNote": payload.get("decisionNote")},
    )


@requires_idempotency
class PlatformReferenceTableListView(APIView):
    """GET: as tabelas (versão em vigor e pendentes). POST: propor uma versão
    nova de uma tabela (nova ou existente), à espera de outra conta."""

    permission_classes = [IsPlatformAdmin]

    def get(self, request):
        return Response(reference_service.list_tables())

    def post(self, request):
        body = ReferenceTableProposalSerializer(data=request.data)
        body.is_valid(raise_exception=True)
        data = body.validated_data
        created = reference_service.propose(
            code=data["code"], name=data["name"], description=data["description"], key_column=data["keyColumn"],
            columns=data["columns"], rows=data["rows"], source=data["source"], effective_from=data["effectiveFrom"],
            reason=data["reason"], actor_id=_actor_id(request),
        )
        _table_audit(request, "platform_reference_table.proposed", created)
        return Response(created, status=status.HTTP_201_CREATED)


class PlatformReferenceTableImportTemplateView(APIView):
    permission_classes = [IsPlatformAdmin]

    def get(self, request):
        from openpyxl import Workbook

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Tabela de referência"
        sheet.append(["Código", "Valor"])
        sheet.append(["LUANDA", 1])
        instructions = workbook.create_sheet("Instruções")
        instructions.append(["Campo", "Formato e valores permitidos"])
        instructions.append(["Código", "Texto único por linha; pode ser usado como coluna-chave."])
        instructions.append(["Valor", "Texto ou número. Acrescente colunas próprias quando necessário."])
        out = BytesIO()
        workbook.save(out)
        response = HttpResponse(out.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        response["Content-Disposition"] = 'attachment; filename="modelo-tabela-referencia.xlsx"'
        return response


class PlatformReferenceTableImportPreviewView(APIView):
    permission_classes = [IsPlatformAdmin]
    parser_classes = [MultiPartParser]

    def post(self, request):
        uploaded = request.FILES.get("file")
        if uploaded is None or not uploaded.name.lower().endswith(".xlsx"):
            return Response({"detail": "Carregue um ficheiro Excel .xlsx."}, status=400)
        if uploaded.size > 10 * 1024 * 1024:
            return Response({"detail": "O ficheiro Excel não pode exceder 10 MB."}, status=400)
        from openpyxl import load_workbook

        try:
            sheet = load_workbook(uploaded, read_only=True, data_only=True).worksheets[0]
            iterator = iter(sheet.iter_rows(values_only=True))
            first = next(iterator, ())
            headers = [str(value or "").strip() for value in first]
            data = []
            for row in iterator:
                if any(value not in (None, "") for value in row):
                    data.append(["" if value is None else str(value).strip() for value in row[:len(headers)]])
                    if len(data) >= 5001:
                        break
        except Exception:
            return Response({"detail": "Não foi possível ler o ficheiro Excel. Verifique se está válido."}, status=400)
        if not headers or not any(headers) or not data:
            return Response({"detail": "O Excel tem de conter cabeçalhos e pelo menos uma linha de dados."}, status=400)
        return Response({"fileName": uploaded.name, "headers": headers, "rows": data[:5000], "truncated": len(data) > 5000})


class PlatformReferenceTableDetailView(APIView):
    permission_classes = [IsPlatformAdmin]

    def get(self, request, code):
        return Response(reference_service.describe(code))


class PlatformReferenceTableVersionsView(APIView):
    permission_classes = [IsPlatformAdmin]

    def get(self, request, code):
        return Response(reference_service.list_versions(code))


@requires_idempotency
class PlatformReferenceTableApproveView(APIView):
    permission_classes = [IsPlatformAdmin]

    def post(self, request, code, version_no):
        body = PlatformParameterDecisionSerializer(data=request.data)
        body.is_valid(raise_exception=True)
        approved = reference_service.approve(code=code, version_no=version_no, actor_id=_actor_id(request),
                                             note=body.validated_data["note"])
        _table_audit(request, "platform_reference_table.approved", approved)
        return Response(approved)


@requires_idempotency
class PlatformReferenceTableRejectView(APIView):
    permission_classes = [IsPlatformAdmin]

    def post(self, request, code, version_no):
        body = PlatformParameterDecisionSerializer(data=request.data)
        body.is_valid(raise_exception=True)
        rejected = reference_service.reject(code=code, version_no=version_no, actor_id=_actor_id(request),
                                            note=body.validated_data["note"])
        _table_audit(request, "platform_reference_table.rejected", rejected)
        return Response(rejected)


# --------------------------------------------------------- modelos de partida
class TemplateProposalSerializer(serializers.Serializer):
    code = serializers.CharField(max_length=60)
    kind = serializers.CharField(max_length=30)
    name = serializers.CharField(max_length=200)
    description = serializers.CharField(required=False, allow_null=True, allow_blank=True, default=None)
    segment = serializers.CharField(required=False, allow_null=True, default=None, max_length=20)
    content = serializers.JSONField()
    source = serializers.CharField(required=False, allow_null=True, allow_blank=True, default=None, max_length=1000)
    reason = serializers.CharField(required=False, allow_null=True, allow_blank=True, default=None)


def _template_audit(request, action: str, payload: dict) -> None:
    record_platform_audit(
        actor_id=_actor_id(request), action=action, entity_type="PlatformTemplate", entity_id=payload.get("code"),
        correlation_id=get_correlation_id(),
        metadata={"versionNo": payload.get("versionNo"), "kind": payload.get("kind"),
                  "configKeys": payload.get("configKeys"), "reason": payload.get("reason"),
                  "decisionNote": payload.get("decisionNote")},
    )


@requires_idempotency
class PlatformTemplateListView(APIView):
    """GET: os modelos (publicado e pendentes). POST: propor uma versão nova
    de um modelo, à espera de outra conta."""

    permission_classes = [IsPlatformAdmin]

    def get(self, request):
        from apps.business_config.application import platform_templates

        return Response(platform_templates.list_all())

    def post(self, request):
        from apps.business_config.application import platform_templates

        body = TemplateProposalSerializer(data=request.data)
        body.is_valid(raise_exception=True)
        data = body.validated_data
        created = platform_templates.propose(
            code=data["code"], kind=data["kind"], name=data["name"], description=data["description"],
            segment=data["segment"], content=data["content"], source=data["source"], reason=data["reason"],
            actor_id=_actor_id(request),
        )
        _template_audit(request, "platform_template.proposed", created)
        return Response(created, status=status.HTTP_201_CREATED)


class PlatformTemplateDetailView(APIView):
    permission_classes = [IsPlatformAdmin]

    def get(self, request, code):
        from apps.business_config.application import platform_templates

        return Response(platform_templates.describe(code))


class PlatformTemplateVersionsView(APIView):
    permission_classes = [IsPlatformAdmin]

    def get(self, request, code):
        from apps.business_config.application import platform_templates

        return Response(platform_templates.list_versions(code))


@requires_idempotency
class PlatformTemplateApproveView(APIView):
    permission_classes = [IsPlatformAdmin]

    def post(self, request, code, version_no):
        from apps.business_config.application import platform_templates

        body = PlatformParameterDecisionSerializer(data=request.data)
        body.is_valid(raise_exception=True)
        approved = platform_templates.approve(code=code, version_no=version_no, actor_id=_actor_id(request),
                                              note=body.validated_data["note"])
        _template_audit(request, "platform_template.approved", approved)
        return Response(approved)


@requires_idempotency
class PlatformTemplateRejectView(APIView):
    permission_classes = [IsPlatformAdmin]

    def post(self, request, code, version_no):
        from apps.business_config.application import platform_templates

        body = PlatformParameterDecisionSerializer(data=request.data)
        body.is_valid(raise_exception=True)
        rejected = platform_templates.reject(code=code, version_no=version_no, actor_id=_actor_id(request),
                                             note=body.validated_data["note"])
        _template_audit(request, "platform_template.rejected", rejected)
        return Response(rejected)
