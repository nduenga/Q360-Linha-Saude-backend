"""Tabelas de referência do fornecedor (camada geral, sem tenant).

Uma tabela tem um código (`provincias`), colunas (`NUMBER` ou `TEXT`), uma
coluna-chave e linhas. Cada alteração é uma versão nova da tabela inteira:
proposta por uma conta da plataforma, aprovada por outra, com vigência. A
versão em vigor é a PUBLISHED de maior número cuja vigência já começou.

Os tenants consultam-nas (só leitura) e usam-nas nas fórmulas com
`tabela("codigo", chave, "coluna")`."""
from __future__ import annotations

import datetime
import re
from decimal import Decimal, InvalidOperation
from typing import Any

from django.db import transaction
from django.db.models import Max
from django.utils import timezone

from apps.foundation.domain.platform_inputs import ReferenceTable, normalise_key
from apps.foundation.application.change_audit import require_reason
from apps.foundation.application.platform_parameter_service import (
    PlatformFourEyesError,
    PlatformParameterStateError,
    UnknownPlatformParameterError,
)
from apps.foundation.domain.system_parameters import InvalidSystemParameterError
from apps.foundation.infrastructure.models import PlatformReferenceTableVersion as Version

COLUMN_TYPES = ("NUMBER", "TEXT")
MAX_COLUMNS = 20
MAX_ROWS = 5000
# Mesmo limite usado em actuarial_import.py — sem isto, um ficheiro enorme
# seria lido inteiro para memória antes do openpyxl sequer começar a parsear.
MAX_WORKBOOK_BYTES = 10 * 1024 * 1024
_CODE = re.compile(r"^[a-z][a-z0-9_]{0,59}$")


def _invalid(message: str) -> InvalidSystemParameterError:
    return InvalidSystemParameterError(message)


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime.datetime, datetime.date)):
        return value.isoformat()
    return str(value).strip()


def parse_workbook_preview(content: bytes) -> dict[str, Any]:
    """Lê a primeira folha de um .xlsx enviado pelo browser e devolve
    cabeçalhos + linhas em texto, prontos para a pré-visualização e para a
    proposta em si (o frontend usa esta resposta diretamente como
    `columns`/`rows`, não só como mostra). Mesma leitura (`openpyxl`,
    `read_only`, `data_only`) já usada em `claims/application/import_service.py`
    para não introduzir uma segunda forma de ler Excel no projecto."""
    import io
    import zipfile

    from openpyxl import load_workbook
    from openpyxl.utils.exceptions import InvalidFileException

    try:
        sheet = load_workbook(io.BytesIO(content), read_only=True, data_only=True).active
    except (InvalidFileException, OSError, KeyError, zipfile.BadZipFile) as exc:
        # BadZipFile (um .xlsx é um zip) e KeyError (zip válido mas sem as
        # partes que um workbook exige) são tão comuns como
        # InvalidFileException num ficheiro mal escolhido pelo utilizador
        # (ex.: um .csv ou .xls antigo com a extensão trocada) — sem isto
        # davam 500 em vez de uma mensagem clara.
        raise _invalid("O ficheiro não é um .xlsx válido.") from exc

    iterator = sheet.iter_rows(values_only=True)
    header_row = next(iterator, None)
    if header_row is None:
        raise _invalid("O ficheiro está vazio.")
    headers = [_cell_text(cell) for cell in header_row]
    headers = headers[: next((i for i in range(len(headers), 0, -1) if headers[i - 1]), 0)]
    if not headers or any(not h for h in headers):
        raise _invalid("A primeira linha tem de ter um nome em cada coluna, sem células vazias.")
    if len(headers) > MAX_COLUMNS:
        raise _invalid(f"Máximo de {MAX_COLUMNS} colunas (o ficheiro tem {len(headers)}).")

    rows: list[list[str]] = []
    truncated = False
    width = len(headers)
    for raw_row in iterator:
        if raw_row is None or not any(cell not in (None, "") for cell in raw_row):
            continue
        if len(rows) >= MAX_ROWS:
            truncated = True
            break
        cells = [_cell_text(raw_row[i]) if i < len(raw_row) else "" for i in range(width)]
        rows.append(cells)

    if not rows:
        raise _invalid("O ficheiro não tem linhas de dados.")

    return {"headers": headers, "rows": rows, "truncated": truncated}


# Dados de teste: fator regional fictício por província.
_REFERENCE_EXAMPLE_ROWS = [
    ["LUA", "Luanda", 1.10], ["BGO", "Bengo", 1.00], ["BEN", "Benguela", 1.00], ["BIE", "Bié", 0.95],
    ["CAB", "Cabinda", 1.05], ["CCU", "Cuando Cubango", 0.95], ["CNO", "Cuanza Norte", 0.95],
    ["CUS", "Cuanza Sul", 0.95], ["CNN", "Cunene", 0.95], ["HUA", "Huambo", 1.00], ["HUI", "Huíla", 1.00],
    ["LNO", "Lunda Norte", 0.95], ["LSU", "Lunda Sul", 0.95], ["MAL", "Malanje", 0.95],
    ["MOX", "Moxico", 0.90], ["NAM", "Namibe", 1.00], ["UIG", "Uíge", 0.95], ["ZAI", "Zaire", 1.00],
]


def build_import_template() -> bytes:
    """Gera, em memória, um .xlsx de exemplo para a importação — não há
    ficheiro estático para manter sincronizado, e a estrutura esperada é
    sempre a mesma (cabeçalhos na primeira linha, uma coluna-chave)."""
    import io

    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Tabela de referência"
    sheet.append(["codigo", "nome", "fator"])
    for row in _REFERENCE_EXAMPLE_ROWS:
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def validate_table(*, code: str, name: str, key_column: str, columns: Any, rows: Any) -> tuple[list[dict], list[dict]]:
    if not isinstance(code, str) or not _CODE.match(code):
        raise _invalid("O código da tabela usa minúsculas, números e '_' e começa por uma letra (ex.: provincias)")
    if not isinstance(name, str) or not name.strip():
        raise _invalid("Indique o nome da tabela")
    if not isinstance(columns, list) or not columns or len(columns) > MAX_COLUMNS:
        raise _invalid(f"Indique entre 1 e {MAX_COLUMNS} colunas")
    clean_columns: list[dict] = []
    seen: set[str] = set()
    for index, col in enumerate(columns):
        if not isinstance(col, dict):
            raise _invalid(f"columns[{index}]: deve ser um objeto")
        col_code = col.get("code")
        if not isinstance(col_code, str) or not _CODE.match(col_code):
            raise _invalid(f"columns[{index}].code: minúsculas, números e '_' (ex.: fator)")
        if col_code in seen:
            raise _invalid(f"columns[{index}].code: '{col_code}' repetido")
        seen.add(col_code)
        col_type = col.get("type", "TEXT")
        if col_type not in COLUMN_TYPES:
            raise _invalid(f"columns[{index}].type: use {list(COLUMN_TYPES)}")
        clean_columns.append({"code": col_code, "label": str(col.get("label") or col_code).strip(), "type": col_type})
    if key_column not in seen:
        raise _invalid(f"A coluna-chave '{key_column}' tem de ser uma das colunas")
    if not isinstance(rows, list) or not rows or len(rows) > MAX_ROWS:
        raise _invalid(f"Indique entre 1 e {MAX_ROWS} linhas")
    types = {c["code"]: c["type"] for c in clean_columns}
    keys: set[str] = set()
    clean_rows: list[dict] = []
    for index, row in enumerate(rows):
        if not isinstance(row, dict):
            raise _invalid(f"rows[{index}]: deve ser um objeto")
        unknown = set(row) - set(types)
        if unknown:
            raise _invalid(f"rows[{index}]: colunas desconhecidas {sorted(unknown)}")
        clean: dict[str, Any] = {}
        for col_code, col_type in types.items():
            cell = row.get(col_code)
            if cell is None or cell == "":
                clean[col_code] = None
                continue
            if col_type == "NUMBER":
                try:
                    number = Decimal(str(cell).strip().replace(",", "."))
                except InvalidOperation:
                    raise _invalid(f"rows[{index}].{col_code}: '{cell}' não é um número") from None
                if not number.is_finite():
                    raise _invalid(f"rows[{index}].{col_code}: número inválido")
                clean[col_code] = format(number.normalize(), "f")
            else:
                clean[col_code] = str(cell).strip()
        key = clean.get(key_column)
        if key is None:
            raise _invalid(f"rows[{index}]: a chave '{key_column}' é obrigatória")
        normal = normalise_key(key)
        if normal in keys:
            raise _invalid(f"rows[{index}]: a chave '{key}' está repetida")
        keys.add(normal)
        clean_rows.append(clean)
    return clean_columns, clean_rows


def effective_version(code: str, *, at: datetime.datetime | None = None) -> Version | None:
    return (
        Version.objects.filter(code=code, status=Version.PUBLISHED, effective_from__lte=at or timezone.now())
        .order_by("-version_no")
        .first()
    )


def effective_tables() -> dict[str, Version]:
    now = timezone.now()
    result: dict[str, Version] = {}
    for version in Version.objects.filter(status=Version.PUBLISHED, effective_from__lte=now).order_by("code", "version_no"):
        result[version.code] = version  # a última (maior número) fica
    return result


def as_reference_table(version: Version) -> ReferenceTable:
    columns = tuple(c["code"] for c in version.columns_json)
    rows = {normalise_key(r[version.key_column]): r for r in version.rows_json if r.get(version.key_column) is not None}
    return ReferenceTable(code=version.code, version_no=version.version_no, key_column=version.key_column,
                          columns=columns, rows=rows)


def snapshot() -> dict[str, int]:
    """{código: versão em vigor} — entra no rasto e na impressão digital das cotações."""
    return {code: v.version_no for code, v in effective_tables().items()}


def version_payload(version: Version, *, in_effect_no: int | None = None, with_rows: bool = True) -> dict[str, Any]:
    payload = {
        "code": version.code,
        "versionNo": version.version_no,
        "status": version.status,
        "name": version.name,
        "description": version.description,
        "keyColumn": version.key_column,
        "columns": version.columns_json,
        "rowCount": len(version.rows_json or []),
        "source": version.source,
        "effectiveFrom": version.effective_from.isoformat() if version.effective_from else None,
        "inEffect": in_effect_no is not None and version.version_no == in_effect_no,
        "reason": version.reason,
        "createdBy": version.created_by,
        "createdAt": version.created_at.isoformat() if version.created_at else None,
        "decidedBy": version.decided_by,
        "decidedAt": version.decided_at.isoformat() if version.decided_at else None,
        "decisionNote": version.decision_note,
    }
    if with_rows:
        payload["rows"] = version.rows_json
    return payload


def list_tables() -> list[dict[str, Any]]:
    """Uma linha por código: a versão em vigor (sem as linhas) e as pendentes."""
    codes = sorted(set(Version.objects.values_list("code", flat=True)))
    result = []
    for code in codes:
        current = effective_version(code)
        pending = Version.objects.filter(code=code, status=Version.PENDING_APPROVAL).order_by("version_no")
        latest = current or Version.objects.filter(code=code).order_by("-version_no").first()
        result.append({
            "code": code,
            "name": latest.name,
            "inEffect": version_payload(current, in_effect_no=current.version_no, with_rows=False) if current else None,
            "pendingVersions": [version_payload(p, with_rows=False) for p in pending],
        })
    return result


def describe(code: str) -> dict[str, Any]:
    current = effective_version(code)
    if current is None and not Version.objects.filter(code=code).exists():
        raise UnknownPlatformParameterError(f"Tabela de referência '{code}' desconhecida")
    return {
        "code": code,
        "inEffect": version_payload(current, in_effect_no=current.version_no) if current else None,
        "pendingVersions": [
            version_payload(p) for p in Version.objects.filter(code=code, status=Version.PENDING_APPROVAL)
        ],
    }


def list_versions(code: str) -> list[dict[str, Any]]:
    current = effective_version(code)
    in_effect_no = current.version_no if current else None
    versions = Version.objects.filter(code=code).order_by("-version_no")
    if not versions.exists():
        raise UnknownPlatformParameterError(f"Tabela de referência '{code}' desconhecida")
    return [version_payload(v, in_effect_no=in_effect_no, with_rows=False) for v in versions]


def propose(
    *, code: str, name: str, key_column: str, columns: Any, rows: Any, reason: str | None, actor_id: str,
    description: str | None = None, source: str | None = None, effective_from: datetime.datetime | None = None,
) -> dict[str, Any]:
    reason = require_reason(reason)
    clean_columns, clean_rows = validate_table(code=code, name=name, key_column=key_column, columns=columns, rows=rows)
    if effective_from is not None and effective_from < timezone.now() - datetime.timedelta(minutes=1):
        raise _invalid("A vigência ('effectiveFrom') não pode estar no passado")
    with transaction.atomic():
        list(Version.objects.select_for_update().filter(code=code).values_list("id", flat=True))
        if Version.objects.filter(code=code, status=Version.PENDING_APPROVAL).exists():
            raise PlatformParameterStateError(f"A tabela '{code}' já tem uma versão à espera de aprovação")
        last = Version.objects.filter(code=code).aggregate(m=Max("version_no"))["m"] or 0
        version = Version.objects.create(
            code=code, version_no=last + 1, name=name.strip(), description=(description or "").strip() or None,
            key_column=key_column, columns_json=clean_columns, rows_json=clean_rows,
            source=(source or "").strip() or None, effective_from=effective_from, reason=reason, created_by=actor_id,
        )
    return version_payload(version)


def _pending(code: str, version_no: int) -> Version:
    version = Version.objects.select_for_update().filter(code=code, version_no=version_no).first()
    if version is None:
        raise UnknownPlatformParameterError(f"A tabela '{code}' não tem a versão {version_no}")
    if version.status != Version.PENDING_APPROVAL:
        raise PlatformParameterStateError(f"A versão {version_no} da tabela '{code}' já foi decidida ({version.status})")
    return version


def approve(*, code: str, version_no: int, actor_id: str, note: str | None = None) -> dict[str, Any]:
    with transaction.atomic():
        version = _pending(code, version_no)
        if version.created_by == actor_id:
            raise PlatformFourEyesError("Quem propôs a versão não a pode aprovar: tem de ser outra conta da plataforma")
        now = timezone.now()
        version.status = Version.PUBLISHED
        version.decided_by = actor_id
        version.decided_at = now
        version.decision_note = (note or "").strip() or None
        if version.effective_from is None or version.effective_from < now:
            version.effective_from = now
        version.save(update_fields=["status", "decided_by", "decided_at", "decision_note", "effective_from"])
    return version_payload(version, with_rows=False)


def reject(*, code: str, version_no: int, actor_id: str, note: str | None) -> dict[str, Any]:
    note = require_reason(note)
    with transaction.atomic():
        version = _pending(code, version_no)
        version.status = Version.REJECTED
        version.decided_by = actor_id if actor_id != version.created_by else None
        version.decided_at = timezone.now()
        version.decision_note = note
        version.save(update_fields=["status", "decided_by", "decided_at", "decision_note"])
    return version_payload(version, with_rows=False)
