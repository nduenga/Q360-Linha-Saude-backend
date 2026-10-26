"""Importação de sinistros e de prémios (R2), apenas Excel (.xlsx).

O formato exato dos dados históricos do cliente ainda é uma decisão em
aberto (OPEN_DECISIONS); este é o formato **documentado** que o simulador
aceita hoje — colunas por nome, qualquer ordem:

  CLAIMS    occurrence_date*, benefit_code*, paid_amount*, member_reference,
            provision_amount, currency (AOA|USD, por omissão AOA),
            status (PAID|OPEN|REJECTED, por omissão PAID)
  PREMIUMS  period_start*, period_end*, premium_amount*, insured_count,
            currency

Datas ISO (`2026-03-31`) ou `31/03/2026`; valores com ponto ou vírgula
decimal. Uma linha inválida nunca aborta as outras: é reportada e saltada.
Serve também para o histórico de 3 anos pedido na transferência de outra
seguradora (`source=OTHER_INSURER`).
"""
from __future__ import annotations

import datetime
import io
from decimal import Decimal, InvalidOperation
from typing import Any

from apps.claims.infrastructure.models import ClaimRecord, ClaimsImport, PremiumRecord
from apps.client.infrastructure.models import Proposer
from apps.foundation.application.audit_service import AuditService
from apps.foundation.application.change_audit import change_metadata
from apps.foundation.application.context import CurrentActorProvider, CurrentTenantProvider
from apps.foundation.domain.problem import DomainProblem
from apps.foundation.infrastructure.audit_repository import DjangoAuditRepositoryAdapter
from apps.product.application.benefit_service import ensure_market_catalog
from apps.product.infrastructure.models import Benefit

_audit = AuditService(DjangoAuditRepositoryAdapter())
KINDS = ("CLAIMS", "PREMIUMS")
SOURCES = ("INTERNAL", "OTHER_INSURER")
_CURRENCIES = ("AOA", "USD")
_STATUSES = ("PAID", "OPEN", "REJECTED")
_REQUIRED = {
    "CLAIMS": ("occurrence_date", "benefit_code", "paid_amount"),
    "PREMIUMS": ("period_start", "period_end", "premium_amount"),
}


class ClaimsImportError(DomainProblem):
    http_status = 400
    problem_code = "VALIDATION-400"


class ClaimsImportNotFoundError(DomainProblem):
    http_status = 404
    problem_code = "NOT-FOUND"


def xlsx_from_rows(rows: list[list]) -> bytes:
    """Livro Excel (.xlsx) com a 1.ª linha como cabeçalho — o único formato que a
    importação aceita. Usado pelos seeds e testes que geram dados de exemplo."""
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    for row in rows:
        sheet.append([v.isoformat() if isinstance(v, datetime.date) else v for v in row])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _rows(filename: str, content: bytes) -> list[dict[str, str]]:
    name = filename.lower()
    if name.endswith(".xlsx"):
        from openpyxl import load_workbook

        sheet = load_workbook(io.BytesIO(content), read_only=True, data_only=True).active
        iterator = sheet.iter_rows(values_only=True)
        aliases = {
            "data de ocorrência": "occurrence_date", "data de ocorrencia": "occurrence_date",
            "código da cobertura": "benefit_code", "codigo da cobertura": "benefit_code",
            "valor pago": "paid_amount", "referência do membro": "member_reference",
            "referencia do membro": "member_reference", "valor provisionado": "provision_amount",
            "moeda": "currency", "estado": "status", "início do período": "period_start",
            "inicio do período": "period_start", "inicio do periodo": "period_start",
            "fim do período": "period_end", "fim do periodo": "period_end",
            "valor do prémio": "premium_amount", "valor do premio": "premium_amount",
            "número de segurados": "insured_count", "numero de segurados": "insured_count",
        }
        header = [aliases.get(str(h or "").strip().casefold(), str(h or "").strip().lower()) for h in next(iterator, [])]
        out = []
        for row in iterator:
            if not any(v not in (None, "") for v in row):
                continue
            out.append({header[i]: ("" if v is None else (v.date().isoformat() if isinstance(v, datetime.datetime) else str(v))).strip()
                        for i, v in enumerate(row) if i < len(header)})
        return out
    raise ClaimsImportError("Formato não suportado. Carregue o modelo Excel .xlsx.")


def _date(raw: str, field: str) -> datetime.date:
    field_label = {"occurrence_date": "Data de ocorrência", "period_start": "Início do período", "period_end": "Fim do período"}.get(field, field)
    for fmt in ("%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.datetime.strptime(raw, fmt).date()
        except ValueError:
            continue
    raise ValueError(f"{field_label} '{raw}' não é uma data (use AAAA-MM-DD ou DD/MM/AAAA)")


def _amount(raw: str, field: str, *, allow_zero: bool = True) -> Decimal:
    field_label = {"paid_amount": "Valor pago", "provision_amount": "Valor provisionado", "premium_amount": "Valor do prémio"}.get(field, field)
    text = raw.replace(" ", "")
    if "," in text and "." not in text:
        text = text.replace(",", ".")
    try:
        value = Decimal(text)
    except InvalidOperation as exc:
        raise ValueError(f"{field_label} '{raw}' não é um número") from exc
    if not value.is_finite() or value < 0 or (value == 0 and not allow_zero):
        raise ValueError(f"{field_label} tem de ser {'igual ou superior a zero' if allow_zero else 'superior a zero'}")
    return value


def import_file(
    *, kind: str, proposer_id, policy_id, source: str, filename: str, content: bytes, dry_run: bool
) -> dict[str, Any]:
    if kind not in KINDS:
        raise ClaimsImportError(f"O tipo (kind) tem de ser um de {list(KINDS)}")
    if source not in SOURCES:
        raise ClaimsImportError(f"A origem (source) tem de ser uma de {list(SOURCES)}")
    proposer = Proposer.objects.filter(id=proposer_id).first()
    if proposer is None:
        raise ClaimsImportError(f"O cliente (proposerId) {proposer_id} não existe neste tenant")
    rows = _rows(filename, content)
    if not rows:
        raise ClaimsImportError("O ficheiro não tem linhas de dados")
    field_labels = {
        "occurrence_date": "Data de ocorrência", "benefit_code": "Código da cobertura", "paid_amount": "Valor pago",
        "period_start": "Início do período", "period_end": "Fim do período", "premium_amount": "Valor do prémio",
    }
    missing = [field_labels[c] for c in _REQUIRED[kind] if c not in rows[0]]
    if missing:
        raise ClaimsImportError(f"Faltam campos obrigatórios no Excel: {', '.join(missing)}")

    tenant_id = CurrentTenantProvider.require().value
    ensure_market_catalog(tenant_id)
    known_benefits = set(Benefit.objects.values_list("code", flat=True))
    valid: list[dict[str, Any]] = []
    errors: list[dict[str, Any]] = []
    for index, row in enumerate(rows, start=2):
        try:
            currency = (row.get("currency") or "AOA").upper()
            if currency not in _CURRENCIES:
                raise ValueError("A moeda tem de ser AOA ou USD")
            if kind == "CLAIMS":
                code = (row.get("benefit_code") or "").upper()
                if code not in known_benefits:
                    raise ValueError(f"O código '{code}' não corresponde a uma cobertura disponível")
                status = (row.get("status") or "PAGO").upper()
                status = {"PAGO": "PAID", "ABERTO": "OPEN", "REJEITADO": "REJECTED"}.get(status, status)
                if status not in _STATUSES:
                    raise ValueError("O Estado tem de ser PAGO, ABERTO ou REJEITADO")
                valid.append(
                    {
                        "occurrence_date": _date(row["occurrence_date"], "occurrence_date"),
                        "benefit_code": code,
                        "paid_amount": _amount(row["paid_amount"], "paid_amount"),
                        "provision_amount": _amount(row.get("provision_amount") or "0", "provision_amount"),
                        "member_reference": row.get("member_reference") or None,
                        "currency_code": currency,
                        "status": status,
                    }
                )
            else:
                start, end = _date(row["period_start"], "period_start"), _date(row["period_end"], "period_end")
                if end < start:
                    raise ValueError("period_end é anterior a period_start")
                insured = row.get("insured_count")
                if insured:
                    if not insured.isdigit() or int(insured) < 1:
                        raise ValueError("insured_count tem de ser um número inteiro positivo")
                valid.append(
                    {
                        "period_start": start,
                        "period_end": end,
                        "premium_amount": _amount(row["premium_amount"], "premium_amount"),
                        "insured_count": int(insured) if insured else None,
                        "currency_code": currency,
                    }
                )
        except (ValueError, KeyError) as exc:
            errors.append({"row": index, "message": str(exc)})

    result: dict[str, Any] = {
        "dryRun": dry_run, "kind": kind, "totalRows": len(rows), "importedRows": len(valid), "errorRows": len(errors), "errors": errors,
    }
    if dry_run:
        return result

    actor = CurrentActorProvider.get()
    job = ClaimsImport.objects.create(
        tenant_id=tenant_id, kind=kind, proposer=proposer, policy_id=policy_id, source=source, file_name=filename,
        total_rows=len(rows), imported_rows=len(valid), error_rows=len(errors), errors_json=errors[:200],
        created_by=actor.actor_id if actor else "unknown",
    )
    model = ClaimRecord if kind == "CLAIMS" else PremiumRecord
    model.objects.bulk_create(
        [model(tenant_id=tenant_id, import_job=job, proposer=proposer, policy_id=policy_id, **item) for item in valid]
    )
    _audit.record(
        action="claims.imported", entity_type="ClaimsImport", entity_id=str(job.id),
        metadata=change_metadata(after={"kind": kind, "importedRows": len(valid), "errorRows": len(errors), "source": source},
                                 proposerId=str(proposer.id)),
    )
    return {**result, "importId": str(job.id)}


def list_imports(*, proposer_id=None, policy_id=None):
    queryset = ClaimsImport.objects.all().order_by("-created_at")
    if proposer_id:
        queryset = queryset.filter(proposer_id=proposer_id)
    if policy_id:
        queryset = queryset.filter(policy_id=policy_id)
    return queryset


def get_import(import_id) -> ClaimsImport:
    job = ClaimsImport.objects.filter(id=import_id).first()
    if job is None:
        raise ClaimsImportNotFoundError(f"A importação de sinistros {import_id} não existe")
    return job


IMPORT_FORMATS: dict[str, dict[str, Any]] = {
    "CLAIMS": {
        "fileFormat": ".xlsx",
        "required": ["Data de ocorrência", "Código da cobertura", "Valor pago"],
        "optional": ["Referência do membro", "Valor provisionado", "Moeda", "Estado"],
        "allowedValues": {"Moeda": ["AOA", "USD"], "Estado": ["PAGO", "ABERTO", "REJEITADO"]},
        "example": ["2026-03-31", "HOSPITALIZATION", "150000,00", "PESS-001", "0", "AOA", "PAGO"],
    },
    "PREMIUMS": {
        "fileFormat": ".xlsx",
        "required": ["Início do período", "Fim do período", "Valor do prémio"],
        "optional": ["Número de segurados", "Moeda"],
        "allowedValues": {"Moeda": ["AOA", "USD"]},
        "example": ["2026-01-01", "2026-12-31", "1200000,00", "25", "AOA"],
    },
}
