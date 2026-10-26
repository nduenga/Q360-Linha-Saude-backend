"""Importação do modelo atuarial a partir do Excel.

Lê o ficheiro de referência (`WORKBOOK_PATH`) ou um ficheiro enviado pela
interface com a mesma estrutura. Os valores calculados (prémios puros,
ajustes adotados) vêm da cache de valores do próprio Excel: o ficheiro tem
de ter sido guardado pelo Excel depois de recalculado.
"""
from __future__ import annotations

import copy
import io
import zipfile
from datetime import date, datetime
from pathlib import Path
from typing import Any, BinaryIO

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from apps.business_config.domain.actuarial_model import RECOMMENDED_CALCULATION_SETTINGS


WORKBOOK_PATH = Path(__file__).with_name("Modelo_Actuarial_Saude_SAU_SAG_Angola.xlsx")
REQUIRED_SHEETS = ("Dashboard_Premissas", "Matriz_SAU", "Matriz_SAG", "Regras_Ajustes")


class ActuarialWorkbookError(ValueError):
    """O ficheiro não é um Excel com a estrutura do modelo atuarial."""
AGE_GENDER_KEYS = ("M|0-17", "F|0-17", "M|18-35", "F|18-35", "M|36-50", "F|36-50", "M|51-65", "F|51-65", "M|>65", "F|>65")
PLAN_CODES = {
    "IND-BASE": ("SAU", "Base"),
    "IND-MEDIO": ("SAU", "Conforto"),
    "IND-SUPERIOR": ("SAU", "Premium"),
    "GRP-COLABORADORES": ("SAG", "Colaboradores"),
    "GRP-QUADROS": ("SAG", "Quadros Médios"),
    "GRP-ADMINISTRACAO": ("SAG", "Administração"),
}
COVERAGE_CODE_MAPPING = {
    "HOSPITALIZATION": ["HOSP_INT", "HOSP_UCI"],
    "AMBULATORY_SURGERY": ["HOSP_CIR"],
    "EMERGENCY": ["URG", "AMBUL"],
    "OUTPATIENT": ["AMB_CONS", "AMB_MCDT"],
    "ENDEMIC_DISEASES": ["MAL"],
    "MATERNITY": ["MAT"],
    "PHARMACY": ["MED"],
    "DENTAL": ["DENT"],
    "PROSTHESIS": ["PROT"],
    "OPTICAL": ["OPT"],
    "CHECKUP_PREVENTIVE": ["PREV"],
    "MEDICAL_EVACUATION": ["EVAC"],
    "DIALYSIS_ONCOLOGY": ["ONCO", "HEMO"],
}


def _value(cell: Any) -> Any:
    value = cell.value
    if isinstance(value, str) and value.startswith("="):
        return None
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return value


def _table(sheet, rows: range, columns: dict[str, int]) -> list[dict[str, Any]]:
    result = []
    for row_number in rows:
        row = {key: _value(sheet.cell(row_number, column)) for key, column in columns.items()}
        if any(value is not None for value in row.values()):
            result.append(row)
    return result


def _adopted(sheet, rows: range, *, value_column: int, adopted_column: int) -> list[dict[str, Any]]:
    """Linhas `valor -> ajuste adotado` lidas dos valores calculados pelo Excel."""
    letter = "ABCDEFGHIJ"[adopted_column - 1]
    return [
        {
            "value": sheet.cell(row, value_column).value,
            "adjustment": sheet.cell(row, adopted_column).value,
            "source_cell": f"{sheet.title}!{letter}{row}",
        }
        for row in rows
    ]


def _open(source: Path | bytes | BinaryIO):
    """Duas vistas do mesmo ficheiro: fórmulas e valores calculados em cache."""
    content = source.read_bytes() if isinstance(source, Path) else source if isinstance(source, bytes) else source.read()
    try:
        # acesso aleatório: read_only seria O(n) por célula
        workbook = load_workbook(io.BytesIO(content), data_only=False)
        cached_workbook = load_workbook(io.BytesIO(content), data_only=True)
    except (InvalidFileException, zipfile.BadZipFile, KeyError, OSError, ValueError) as exc:
        raise ActuarialWorkbookError("O ficheiro não é um livro Excel (.xlsx) válido.") from exc
    missing = [name for name in REQUIRED_SHEETS if name not in workbook.sheetnames]
    if missing:
        raise ActuarialWorkbookError(f"Faltam as folhas do modelo atuarial: {', '.join(missing)}.")
    if cached_workbook["Matriz_SAU"].cell(6, 19).value is None:
        raise ActuarialWorkbookError(
            "O ficheiro não tem os valores calculados (Matriz_SAU!S6 está vazio). Abra-o no Excel, "
            "deixe recalcular e guarde-o antes de o importar."
        )
    return workbook, cached_workbook


def build_actuarial_model(source: Path | bytes | BinaryIO = WORKBOOK_PATH, *, source_name: str | None = None) -> dict[str, Any]:
    workbook, cached_workbook = _open(source)
    dashboard = workbook["Dashboard_Premissas"]
    sau = workbook["Matriz_SAU"]
    sag = workbook["Matriz_SAG"]
    group_rules = workbook["Matriz_SAG"]
    adjustments = workbook["Regras_Ajustes"]
    cached_adjustments = cached_workbook["Regras_Ajustes"]

    coverage_assumptions = _table(
        dashboard,
        range(79, 96),
        {
            "name": 2, "code": 3, "group": 4, "frequency_unit": 5, "reference_cell": 6,
            "frequency_per_life": 7, "frequency_per_1000_lives": 8, "mean_network_cost_aoa": 9,
            "mean_reimbursement_cost_aoa": 10, "network_utilization": 11, "weighted_mean_cost_aoa": 12,
            "severity_cv": 13, "apply_network_factor": 14, "technical_rationale": 15,
        },
    )
    age_relativities = []
    for row_number in range(114, 131):
        code = _value(dashboard.cell(row_number, 3))
        if code is not None:
            age_relativities.append({
                "coverage_code": code,
                "factors": {key: _value(dashboard.cell(row_number, index)) for index, key in enumerate(AGE_GENDER_KEYS, start=4)},
            })

    plan_coverages = {}
    for segment, formula_sheet_name in (("SAU", "Matriz_SAU"), ("SAG", "Matriz_SAG")):
        formula_sheet = workbook[formula_sheet_name]
        cached_sheet = cached_workbook[formula_sheet_name]
        plan_coverages[segment] = [
            {
                "plan": _value(formula_sheet.cell(row_number, 2)),
                "coverage_code": _value(formula_sheet.cell(row_number, 3)),
                "status": _value(formula_sheet.cell(row_number, 6)),
                "capital_sublimit_aoa": cached_sheet.cell(row_number, 7).value,
            }
            for row_number in range(6, 57)
            if _value(formula_sheet.cell(row_number, 2)) is not None
        ]
    reference_pure_premiums = []
    for segment, formula_sheet_name in (("SAU", "Matriz_SAU"), ("SAG", "Matriz_SAG")):
        formula_sheet = workbook[formula_sheet_name]
        cached_sheet = cached_workbook[formula_sheet_name]
        for row_number in range(6, 57):
            plan = _value(formula_sheet.cell(row_number, 2))
            coverage_code = _value(formula_sheet.cell(row_number, 3))
            status = _value(formula_sheet.cell(row_number, 6))
            pure_premium = cached_sheet.cell(row_number, 19).value
            if not plan or not coverage_code or status == "—" or pure_premium is None:
                continue
            lambda_per_life = cached_sheet.cell(row_number, 8).value
            adjusted_mean_cost = cached_sheet.cell(row_number, 12).value
            trend_factor = cached_sheet.cell(row_number, 13).value
            severity_cv = cached_sheet.cell(row_number, 15).value
            mean_claim = cached_sheet.cell(row_number, 14).value
            sigma_ln = cached_sheet.cell(row_number, 16).value
            mu_ln = cached_sheet.cell(row_number, 17).value
            lev_factor = cached_sheet.cell(row_number, 18).value
            reference_pure_premiums.append({
                "segment": segment,
                "plan": plan,
                "coverage_code": coverage_code,
                "status": status,
                "capital_sublimit_aoa": _value(formula_sheet.cell(row_number, 7)),
                "reference_age_band": "18-35",
                "reference_sex": "M",
                "annual_pure_premium": str(pure_premium),
                "frequency_per_life": str(lambda_per_life),
                "adjusted_mean_cost_aoa": str(adjusted_mean_cost),
                "trend_factor": str(trend_factor),
                "severity_cv": str(severity_cv),
                "mean_claim_aoa": str(mean_claim),
                "sigma_ln": str(sigma_ln),
                "mu_ln": str(mu_ln),
                "lev_factor": str(lev_factor),
                "source_cell": f"{formula_sheet_name}!S{row_number}",
            })

    package_rows = _table(
        adjustments,
        range(49, 54),
        {
            "code": 2, "copayment_aoa": 3, "reimbursement_share": 4,
            "annual_outpatient_deductible_aoa": 5, "inpatient_deductible_aoa": 6,
            "reference_adjustment_sag": 9, "reference_adjustment_sau": 10,
        },
    )

    return {
        "calculation_settings": copy.deepcopy(RECOMMENDED_CALCULATION_SETTINGS),
        "model_status": "DRAFT_REQUIRES_ACTUARIAL_APPROVAL",
        "source_file": source_name or (source.name if isinstance(source, Path) else "upload.xlsx"),
        "source_sheets": ["Dashboard_Premissas", "Matriz_SAU", "Matriz_SAG", "Regras_Ajustes"],
        "plan_code_mapping": {code: {"segment": segment, "actuarial_plan": actuarial_plan} for code, (segment, actuarial_plan) in PLAN_CODES.items()},
        "coverage_code_mapping": COVERAGE_CODE_MAPPING,
        "currency_code": "AOA",
        "currency_label": "Kwanza (Kz)",
        "reference_date": _value(dashboard["D38"]),
        "approval_warning": _value(dashboard["A2"]),
        "assumptions": {
            "usd_to_aoa": _value(dashboard["D39"]),
            "general_inflation": _value(dashboard["D40"]),
            "medical_inflation_excess": _value(dashboard["D41"]),
            "imported_cost_fx_component": _value(dashboard["D42"]),
            "projection_months": _value(dashboard["D44"]),
            "loadings": {
                segment: {
                    "volatility": _value(dashboard.cell(50, column)),
                    "administration_expense": _value(dashboard.cell(51, column)),
                    "acquisition_commission": _value(dashboard.cell(52, column)),
                    "insurer_margin": _value(dashboard.cell(53, column)),
                    "stamp_duty": _value(dashboard.cell(55, column)),
                    "arseg_levy": _value(dashboard.cell(56, column)),
                }
                for segment, column in (("SAU", 4), ("SAG", 5))
            },
            "payment_installments": _table(
                dashboard, range(61, 65), {"frequency": 2, "installments": 3, "loading": 4}
            ),
            "age_exposure": _table(
                dashboard, range(69, 74),
                {"age_band": 2, "age_from": 3, "age_to": 4, "male_lives": 5, "female_lives": 6},
            ),
        },
        "coverage_assumptions": coverage_assumptions,
        "age_gender_relativities": age_relativities,
        "plan_coverages": plan_coverages,
        "reference_pure_premiums": reference_pure_premiums,
        "plan_factors": _table(
            dashboard, range(154, 160),
            {"description": 2, "product_plan": 3, "network_factor": 4, "room_limit_per_day_aoa": 5, "icu_annual_limit_aoa": 6},
        ),
        "group_size_factors": _table(
            group_rules, range(100, 104),
            {"label": 2, "headcount_from": 3, "headcount_to": 4, "premium_adjustment": 5, "volatility": 6, "rationale": 7},
        ),
        "group_participation_factors": _table(
            group_rules, (106, 107, 109, 110, 111, 112),
            {"rule": 2, "minimum_participation": 3, "premium_adjustment": 4, "rationale": 5},
        ),
        "industry_factors": _table(
            group_rules, range(115, 122), {"industry": 2, "factor": 3, "rationale": 4}
        ),
        "adjustment_assumptions": {
            "copayment_per_visit_aoa": [_value(adjustments.cell(row, 2)) for row in range(24, 28)],
            "reimbursement_shares": [_value(adjustments.cell(row, 2)) for row in range(30, 34)],
            "outpatient_deductibles_aoa": [_value(adjustments.cell(row, 2)) for row in range(36, 40)],
            "inpatient_deductibles_aoa": [_value(adjustments.cell(row, 2)) for row in range(42, 46)],
            "participation_packages": package_rows,
            "pre_existing_severity": _table(
                adjustments, range(58, 63), {"severity": 2, "surcharge": 3, "decision": 4}
            ),
            "entry_age_adjustments": _table(
                adjustments, range(79, 84), {"age_from": 2, "surcharge": 3, "decision": 4}
            ),
            "family_size_discounts": _table(
                adjustments, range(88, 92), {"minimum_members": 2, "discount": 3}
            ),
            "cross_sell_discounts": _table(
                adjustments, range(94, 99), {"product": 2, "discount": 3}
            ),
            "no_claims_renewal_discount": _value(adjustments["C100"]),
            "maximum_cumulative_discount": _value(adjustments["C101"]),
            "regional_factors": _table(
                adjustments, range(104, 109), {"region": 2, "factor": 3, "rationale": 4}
            ),
            "copayment_demand_elasticity": _value(adjustments["C15"]),
            "reimbursement_behavioral_elasticity": _value(adjustments["C16"]),
            "annual_outpatient_utilization_probability": _value(adjustments["C18"]),
            # Calculador!C55:C58 — ajuste ADOPTADO (com override aprovado, se houver)
            # por valor escolhido; o Calculador usa MATCH exacto, por isso só os
            # valores da tabela são aceites.
            "cost_sharing": {
                "copayment_per_visit_aoa": _adopted(cached_adjustments, range(24, 28), value_column=2, adopted_column=8),
                "reimbursement_share": _adopted(cached_adjustments, range(30, 34), value_column=2, adopted_column=6),
                "outpatient_annual_deductible_aoa": _adopted(cached_adjustments, range(36, 40), value_column=2, adopted_column=6),
                "inpatient_deductible_aoa": _adopted(cached_adjustments, range(42, 46), value_column=2, adopted_column=6),
            },
        },
    }