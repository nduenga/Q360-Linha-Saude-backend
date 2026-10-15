"""Da estrutura da proposta ao documento: um modelo intermédio (blocos) que
os três renderizadores — HTML, PDF (reportlab) e DOCX (python-docx) —
percorrem. A lógica de cada secção vive só aqui, por isso os três formatos
dizem sempre exactamente o mesmo."""
from __future__ import annotations

import base64
import html
import io
from typing import Any

from apps.business_config.application.branding import load_logo
from apps.quote.domain.proposal_template import SECTIONS, render_text

Block = tuple  # ("title"|"heading"|"subheading"|"paragraph"|"note"|"bullets"|"kv"|"table"|"signature", ...)


def variables_for(content: dict[str, Any], *, document_number: str | None, generated_on: str) -> dict[str, Any]:
    meta, client, plan, premium = content["meta"], content["client"], content["plan"], content["premium"]
    return {
        "quote.number": meta["quoteNumber"],
        "quote.date": meta["quoteDate"],
        "quote.validUntil": meta["validUntil"],
        "quote.type": meta["quoteType"],
        "document.number": document_number or "(a atribuir)",
        "document.date": generated_on,
        "client.name": client["name"],
        "client.nif": client["nif"],
        "plan.name": plan["name"],
        "plan.product": plan["product"],
        "plan.segment": plan["segment"],
        "premium.total": premium["totalText"],
        "premium.currency": premium["currency"],
        "scenario.name": (meta["scenario"] or {}).get("name"),
        "insured.count": len(content["population"]),
    }


def build_blocks(content: dict[str, Any], spec: dict[str, Any], variables: dict[str, Any]) -> list[Block]:
    blocks: list[Block] = []
    branding = content.get("branding")
    if branding and (branding.get("company") or branding.get("color")):
        blocks.append(("brand", {**branding, "logo": load_logo(branding)}))
    blocks += [("title", render_text(spec["title"], variables)), ("paragraph", render_text(spec["intro"], variables))]
    builders = {
        "summary": _summary,
        "population": _population,
        "coverages": _coverages,
        "premium": _premium,
        "waiting": _waiting,
        "preexisting": _preexisting,
        "exclusions": _exclusions,
        "conditions": _conditions,
        "signature": _signature,
    }
    for section in spec["sections"]:
        builders[section](content, blocks)
    for text in content.get("businessTexts") or []:
        blocks += [("heading", text.get("title") or "Informação"), ("paragraph", text["text"])]
    if content.get("requiredDocuments"):
        blocks += [
            ("heading", "Documentos a apresentar"),
            ("bullets", [f"{d['documentType']}" + (f" — {d['description']}" if d.get("description") else "") for d in content["requiredDocuments"]]),
        ]
    disclaimer = content.get("disclaimer") and render_text(spec.get("disclaimer") or content["disclaimer"], variables)
    if disclaimer:
        blocks.append(("note", disclaimer))
    if spec.get("footer"):
        blocks.append(("note", render_text(spec["footer"], variables)))
    if branding and branding.get("footer"):
        blocks.append(("note", branding["footer"]))
    return blocks


def _summary(content, blocks):
    meta, client, plan = content["meta"], content["client"], content["plan"]
    rows = [
        ("Cliente", client["name"]),
        ("NIF", client["nif"] or "—"),
        ("Tipo", meta["quoteType"]),
        ("Produto", plan["product"]),
        ("Plano", plan["name"]),
        ("Segmento", plan["segment"]),
    ]
    if plan["tier"]:
        rows.append(("Escalão", plan["tier"]))
    if plan["category"]:
        rows.append(("Categoria", plan["category"]))
    if plan["accessSystem"]:
        rows.append(("Sistema de acesso", plan["accessSystem"]))
    rows += [("Data da cotação", meta["quoteDate"]), ("Proposta válida até", meta["validUntil"] or "—")]
    if meta["scenario"]:
        rows.append(("Cenário", meta["scenario"]["name"]))
    blocks += [("heading", SECTIONS["summary"].split(" (")[0]), ("kv", rows)]


def _population(content, blocks):
    people = content["population"]
    rows = [
        [
            p["name"], p["relationship"], "—" if p["age"] is None else str(p["age"]),
            "Sim" if p["clinicalQuestionnaireRequired"] else "Não", "Sim" if p["medicalExamRequired"] else "Não",
        ]
        for p in people
    ]
    blocks += [
        ("heading", "Pessoas seguras"),
        ("table", ["Pessoa segura", "Relação", "Idade", "Questionário clínico", "Exames médicos"], rows),
    ]


def _coverages(content, blocks):
    blocks.append(("heading", "Coberturas, capitais e participação"))
    option = content["participationOption"]
    if option:
        rows = [("Opção de participação", f"{option['code']} — {option['name'] or ''}".strip(" —"))]
        if option["annualDeductible"] is not None:
            rows.append(("Franquia anual", option["annualDeductible"]))
        if option["copayPerConsultation"] is not None:
            rows.append(("Copagamento por consulta", option["copayPerConsultation"]))
        if option["reimbursementPct"] is not None:
            rows.append(("Comparticipação no reembolso", f"{option['reimbursementPct']}%"))
        if option["hospitalOnly"]:
            rows.append(("Âmbito", "Só hospitalização (ambulatório não coberto)"))
        blocks.append(("kv", rows))
    contracted = [c for c in content["coverages"] if c["contracted"]]
    for coverage in contracted:
        title = coverage["name"] + (" (opcional escolhida)" if coverage["inclusion"] == "OPTIONAL" else "")
        lines: list[str] = []
        lines += coverage["limits"]
        lines += coverage["participation"]
        if coverage["waiting"]:
            lines.append(f"Carência: {coverage['waiting']}")
        if coverage["requiresPriorApproval"]:
            lines.append("Aprovação prévia obrigatória")
        if coverage["destinations"]:
            lines.append("Destinos: " + ", ".join(coverage["destinations"]))
        if coverage["currency"] != "AOA":
            lines.append(f"Capital em {coverage['currency']} (risco cambial)")
        if coverage["closedDiseaseList"]:
            lines.append("Doenças abrangidas: " + ", ".join(coverage["closedDiseaseList"]))
        if coverage["maxDaysPerYear"]:
            lines.append(f"Máximo de {coverage['maxDaysPerYear']} dias por ano")
        if coverage["timeDeductibleDays"]:
            lines.append(f"Franquia temporal: {coverage['timeDeductibleDays']} dias de internamento")
        if coverage["notes"]:
            lines.append(coverage["notes"])
        blocks.append(("subheading", title))
        blocks.append(("bullets", lines or ["Sem limites definidos nesta proposta"]))
    optional_rest = [c["name"] for c in content["coverages"] if not c["contracted"]]
    if optional_rest:
        blocks += [("subheading", "Coberturas opcionais disponíveis (não incluídas)"), ("bullets", optional_rest)]


def _premium(content, blocks):
    premium = content["premium"]
    blocks.append(("heading", "Prémio"))
    rows = [[m["name"], _money_text(m["total"], premium["currency"])] for m in premium["members"]]
    if rows:
        blocks.append(("table", ["Pessoa segura", "Prémio"], rows))
    if premium["adjustments"]:
        blocks += [("subheading", "Ajustes comerciais"), ("bullets", premium["adjustments"])]
    blocks.append(("kv", [("Prémio total", premium["totalText"])]))
    for warning in premium["warnings"]:
        blocks.append(("note", warning))


def _money_text(amount: str, currency: str) -> str:
    from apps.quote.application.proposal_content import money

    return money(amount, currency)


def _waiting(content, blocks):
    rows = [[w["coverage"], w["text"]] for w in content["waiting"]]
    blocks.append(("heading", "Períodos de carência"))
    blocks.append(("table", ["Cobertura", "Carência"], rows) if rows else ("paragraph", "Sem períodos de carência definidos."))


def _preexisting(content, blocks):
    info = content["preExisting"]
    blocks.append(("heading", "Doenças preexistentes"))
    if info["defaultTreatment"]:
        blocks.append(("paragraph", "Tratamento previsto no plano: " + info["defaultTreatment"]))
    decisions = info["decisions"]
    if decisions:
        rows = []
        for person in decisions:
            decision = person["preExisting"]
            detail = decision["text"]
            if decision["condition"]:
                detail += f" Doença: {decision['condition']}."
            if decision["waitingMonths"]:
                detail += f" Carência: {decision['waitingMonths']} meses."
            if decision["loadingPct"]:
                detail += f" Sobreprémio: {decision['loadingPct']}%."
            if decision["sublimitAmount"]:
                detail += f" Sublimite: {decision['sublimitAmount']}."
            rows.append([person["name"], detail])
        blocks.append(("table", ["Pessoa segura", "Decisão de subscrição"], rows))
    elif not info["defaultTreatment"]:
        blocks.append(("paragraph", "Não há decisões de subscrição sobre doenças preexistentes."))


def _exclusions(content, blocks):
    blocks += [("heading", "Exclusões"), ("bullets", content["exclusions"])]


def _conditions(content, blocks):
    blocks += [("heading", "Condições"), ("bullets", content["conditions"])]


def _signature(content, blocks):
    blocks += [
        ("heading", "Aceitação"),
        ("paragraph", "O tomador declara conhecer e aceitar as condições desta proposta."),
        ("signature", ["Assinatura do tomador", "Data"]),
    ]


# ---------------------------------------------------------------- renderers
def _brand_html(brand: dict[str, Any]) -> str:
    e = html.escape
    color = brand.get("color") or "#1a1a1a"
    logo = ""
    if brand.get("logo"):
        mime = "image/png" if brand["logo"].startswith(bytes([0x89, 0x50, 0x4E, 0x47])) else "image/jpeg"
        logo = f"<img alt='' style='max-height:64px;margin-right:1rem' src='data:{mime};base64,{base64.b64encode(brand['logo']).decode()}'>"
    name = f"<strong style='font-size:1.3rem;color:{e(color)}'>{e(brand['company'])}</strong>" if brand.get("company") else ""
    tagline = f"<div class='note'>{e(brand['tagline'])}</div>" if brand.get("tagline") else ""
    return f"<header style='display:flex;align-items:center;border-bottom:3px solid {e(color)};padding-bottom:.6rem;margin-bottom:1rem'>{logo}<div>{name}{tagline}</div></header>"


def render_html(blocks: list[Block], *, title: str) -> bytes:
    e = html.escape
    out = [
        "<!doctype html><html lang='pt-AO'><head><meta charset='utf-8'>",
        f"<title>{e(title)}</title>",
        "<style>body{font-family:Arial,Helvetica,sans-serif;max-width:860px;margin:2rem auto;color:#1a1a1a;line-height:1.45}"
        "h1{font-size:1.6rem}h2{font-size:1.2rem;border-bottom:1px solid #ccc;padding-bottom:.2rem;margin-top:1.6rem}"
        "h3{font-size:1rem;margin-bottom:.2rem}table{border-collapse:collapse;width:100%;margin:.5rem 0}"
        "td,th{border:1px solid #ccc;padding:.35rem .5rem;text-align:left;vertical-align:top}th{background:#f3f3f3}"
        ".note{font-size:.85rem;color:#555}.sig{margin-top:2rem}.sig span{display:inline-block;width:45%;border-top:1px solid #000;margin-right:4%}"
        "</style></head><body>",
    ]
    for block in blocks:
        kind = block[0]
        if kind == "brand":
            out.append(_brand_html(block[1]))
        elif kind == "title":
            out.append(f"<h1>{e(block[1])}</h1>")
        elif kind == "heading":
            out.append(f"<h2>{e(block[1])}</h2>")
        elif kind == "subheading":
            out.append(f"<h3>{e(block[1])}</h3>")
        elif kind == "paragraph":
            out.append(f"<p>{e(block[1])}</p>")
        elif kind == "note":
            out.append(f"<p class='note'>{e(block[1])}</p>")
        elif kind == "bullets":
            out.append("<ul>" + "".join(f"<li>{e(item)}</li>" for item in block[1]) + "</ul>")
        elif kind == "kv":
            out.append("<table>" + "".join(f"<tr><th>{e(k)}</th><td>{e(str(v))}</td></tr>" for k, v in block[1]) + "</table>")
        elif kind == "table":
            head = "".join(f"<th>{e(h)}</th>" for h in block[1])
            body = "".join("<tr>" + "".join(f"<td>{e(str(c))}</td>" for c in row) + "</tr>" for row in block[2])
            out.append(f"<table><tr>{head}</tr>{body}</table>")
        elif kind == "signature":
            out.append("<div class='sig'>" + "".join(f"<span>{e(label)}</span>" for label in block[1]) + "</div>")
    out.append("</body></html>")
    return "".join(out).encode("utf-8")


def render_pdf(blocks: list[Block], *, title: str, layout: dict[str, Any] | None = None) -> bytes:
    """`layout` (opcional) ajusta a página ao papel timbrado do GTE:
    `{"pageSize": (largura, altura) em pontos, "margins": (esq, dir, topo, base) em pontos}`."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    styles = getSampleStyleSheet()
    body = ParagraphStyle("body", parent=styles["BodyText"], fontSize=9.5, leading=13)
    cell = ParagraphStyle("cell", parent=body, fontSize=8.5, leading=11)
    note = ParagraphStyle("note", parent=body, fontSize=8, textColor=colors.HexColor("#555555"))
    esc = html.escape
    buffer = io.BytesIO()
    page_size = tuple((layout or {}).get("pageSize") or A4)
    left, right, top, bottom = (layout or {}).get("margins") or (2 * cm, 2 * cm, 2 * cm, 2 * cm)
    doc = SimpleDocTemplate(
        buffer, pagesize=page_size, leftMargin=left, rightMargin=right, topMargin=top, bottomMargin=bottom, title=title
    )
    width = page_size[0] - left - right
    story: list[Any] = []

    def table(rows, header=None, col_widths=None, header_bold=True):
        data = [[Paragraph(esc(str(c)), cell) for c in row] for row in rows]
        if header:
            data.insert(0, [Paragraph(f"<b>{esc(h)}</b>", cell) for h in header])
        t = Table(data, colWidths=col_widths, repeatRows=1 if header else 0)
        style = [("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#bbbbbb")), ("VALIGN", (0, 0), (-1, -1), "TOP")]
        if header:
            style.append(("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#eeeeee")))
        t.setStyle(TableStyle(style))
        return t

    for block in blocks:
        kind = block[0]
        if kind == "brand":
            brand = block[1]
            accent = colors.HexColor(brand["color"]) if brand.get("color") else colors.black
            head_style = ParagraphStyle("brand", parent=styles["Heading2"], textColor=accent)
            cells = []
            if brand.get("logo"):
                from reportlab.platypus import Image

                cells.append(Image(io.BytesIO(brand["logo"]), width=2.2 * cm, height=2.2 * cm, kind="proportional"))
            text = f"<b>{esc(brand['company'])}</b>" if brand.get("company") else ""
            if brand.get("tagline"):
                text += f"<br/><font size=8>{esc(brand['tagline'])}</font>"
            cells.append(Paragraph(text, head_style))
            header = Table([cells], colWidths=[2.6 * cm, width - 2.6 * cm] if len(cells) == 2 else [width])
            header.setStyle(TableStyle([("LINEBELOW", (0, 0), (-1, -1), 2, accent), ("VALIGN", (0, 0), (-1, -1), "MIDDLE")]))
            story += [header, Spacer(1, 10)]
        elif kind == "title":
            story.append(Paragraph(esc(block[1]), styles["Title"]))
        elif kind == "heading":
            story += [Spacer(1, 8), Paragraph(esc(block[1]), styles["Heading2"])]
        elif kind == "subheading":
            story.append(Paragraph(esc(block[1]), styles["Heading4"]))
        elif kind == "paragraph":
            story += [Paragraph(esc(block[1]), body), Spacer(1, 4)]
        elif kind == "note":
            story += [Spacer(1, 4), Paragraph(esc(block[1]), note)]
        elif kind == "bullets":
            story.append(ListFlowable([ListItem(Paragraph(esc(item), body)) for item in block[1]], bulletType="bullet", leftIndent=14))
        elif kind == "kv":
            rows = [[f"<b>{esc(k)}</b>", esc(str(v))] for k, v in block[1]]
            t = Table([[Paragraph(a, cell), Paragraph(b, cell)] for a, b in rows], colWidths=[width * 0.35, width * 0.65])
            t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#bbbbbb")), ("VALIGN", (0, 0), (-1, -1), "TOP")]))
            story.append(t)
        elif kind == "table":
            columns = len(block[1])
            widths = [width / columns] * columns
            story.append(table(block[2], header=block[1], col_widths=widths))
        elif kind == "signature":
            story.append(Spacer(1, 30))
            cells = [Paragraph("_" * 38 + "<br/>" + esc(label), cell) for label in block[1]]
            story.append(Table([cells], colWidths=[width / len(block[1])] * len(block[1])))
    doc.build(story)
    return buffer.getvalue()


def render_docx(blocks: list[Block], *, title: str) -> bytes:
    from docx import Document

    document = Document()
    document.core_properties.title = title
    for block in blocks:
        kind = block[0]
        if kind == "brand":
            brand = block[1]
            if brand.get("logo"):
                from docx.shared import Cm

                document.add_picture(io.BytesIO(brand["logo"]), height=Cm(1.8))
            if brand.get("company"):
                run = document.add_paragraph().add_run(brand["company"])
                run.bold = True
                if brand.get("color"):
                    from docx.shared import RGBColor

                    run.font.color.rgb = RGBColor.from_string(brand["color"][1:].upper())
            if brand.get("tagline"):
                document.add_paragraph(brand["tagline"])
        elif kind == "title":
            document.add_heading(block[1], level=0)
        elif kind == "heading":
            document.add_heading(block[1], level=1)
        elif kind == "subheading":
            document.add_heading(block[1], level=2)
        elif kind in ("paragraph", "note"):
            paragraph = document.add_paragraph(block[1])
            if kind == "note":
                for run in paragraph.runs:
                    run.italic = True
        elif kind == "bullets":
            for item in block[1]:
                document.add_paragraph(item, style="List Bullet")
        elif kind == "kv":
            table = document.add_table(rows=0, cols=2)
            table.style = "Table Grid"
            for key, value in block[1]:
                cells = table.add_row().cells
                cells[0].text, cells[1].text = key, str(value)
        elif kind == "table":
            table = document.add_table(rows=1, cols=len(block[1]))
            table.style = "Table Grid"
            for index, head in enumerate(block[1]):
                table.rows[0].cells[index].text = head
            for row in block[2]:
                cells = table.add_row().cells
                for index, value in enumerate(row):
                    cells[index].text = str(value)
        elif kind == "signature":
            document.add_paragraph()
            table = document.add_table(rows=1, cols=len(block[1]))
            for index, label in enumerate(block[1]):
                table.rows[0].cells[index].text = "_" * 30 + "\n" + label
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


RENDERERS = {
    "HTML": (render_html, "text/html; charset=utf-8", "html"),
    "PDF": (render_pdf, "application/pdf", "pdf"),
    "DOCX": (render_docx, "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "docx"),
}
