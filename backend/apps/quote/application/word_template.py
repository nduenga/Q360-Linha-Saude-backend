"""Modelo Word (opção B): o GTE escreve o documento no Word, com marcas
`{{variável}}` onde quer os dados, e o sistema substitui-as.

Duas espécies de marcas:

* **Valores simples** — `{{client.name}}`, `{{policy.number}}`… trocam-se pelo valor.
* **Listas** — `{{insured.name}}`, `{{coverage.limits}}`, `{{installment.amount}}`…
  Uma **linha de tabela** (ou um **parágrafo**) que tenha uma marca de lista
  repete-se uma vez por elemento: escreva a linha-modelo uma só vez e o sistema
  cria as restantes.

Limites: escreva cada marca de seguida, sem mudar de formato a meio; as marcas
dentro de caixas de texto e formas não são substituídas.
"""
from __future__ import annotations

import copy
import io
import re
from typing import Any, Iterable

from apps.quote.domain.proposal_template import VARIABLES as PROPOSAL_VARIABLES
from apps.quote.domain.proposal_template import InvalidTemplateError

PLACEHOLDER = re.compile(r"\{\{\s*([A-Za-z0-9_.]+)\s*\}\}")
DOCX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

COMMON_SCALARS = {
    "client.name": "Nome do cliente/tomador",
    "client.nif": "NIF ou identificação do cliente",
    "client.email": "Email do cliente",
    "client.phone": "Telefone do cliente",
    "plan.name": "Nome do plano",
    "plan.product": "Nome do produto",
    "plan.segment": "Segmento (Individual e Familiar / Grupo)",
    "plan.accessSystem": "Sistema de acesso (rede / reembolso / misto)",
    "premium.total": "Prémio total (com moeda)",
    "premium.currency": "Moeda",
    "document.number": "Número do documento",
    "document.date": "Data de emissão do documento",
    "insured.count": "Número de pessoas seguras",
}
PROPOSAL_SCALARS = {
    **COMMON_SCALARS,
    **{k: v for k, v in PROPOSAL_VARIABLES.items() if k not in COMMON_SCALARS},
}
POLICY_SCALARS = {
    **COMMON_SCALARS,
    "quote.number": "Número da cotação de origem",
    "policy.number": "Número da apólice",
    "policy.status": "Estado da apólice",
    "policy.type": "Tipo (Individual / Grupo)",
    "policy.issueDate": "Data de emissão da apólice",
    "policy.startDate": "Início de vigência",
    "policy.endDate": "Fim de vigência",
    "policy.termMonths": "Duração em meses",
    "policy.paymentFrequency": "Fracionamento do prémio",
    "policy.totalPremium": "Prémio total anual",
    "signature.name": "Nome de quem assinou",
    "signature.document": "Documento de identificação de quem assinou",
    "signature.capacity": "Qualidade de quem assinou",
    "signature.date": "Data da assinatura",
    "payment.firstDate": "Data do pagamento da 1.ª prestação",
    "payment.firstMethod": "Meio de pagamento da 1.ª prestação",
    "payment.firstReference": "Referência do pagamento da 1.ª prestação",
}
SCALARS_BY_TYPE = {"QUOTE_PROPOSAL": PROPOSAL_SCALARS, "POLICY": POLICY_SCALARS}

GROUPS: dict[str, dict[str, Any]] = {
    "insured": {"label": "Pessoas seguras", "fields": {"name": "Nome", "relationship": "Relação", "age": "Idade", "status": "Estado"}},
    "coverage": {
        "label": "Coberturas contratadas",
        "fields": {"name": "Cobertura", "limits": "Capitais e limites", "participation": "Participação/franquia", "waiting": "Carência", "priorApproval": "Aprovação prévia (Sim/Não)"},
    },
    "installment": {
        "label": "Prestações do prémio (só na apólice)",
        "fields": {"no": "Número", "period": "Período", "dueDate": "Vencimento", "amount": "Valor", "status": "Estado"},
    },
    "exclusion": {"label": "Exclusões", "fields": {"text": "Texto"}},
    "condition": {"label": "Condições", "fields": {"text": "Texto"}},
    "waiting": {"label": "Períodos de carência", "fields": {"coverage": "Cobertura", "text": "Carência"}},
}
GROUP_FIELDS = {f"{g}.{f}" for g, spec in GROUPS.items() for f in spec["fields"]}


def allowed_names(document_type: str) -> set[str]:
    return set(SCALARS_BY_TYPE.get(document_type, PROPOSAL_SCALARS)) | GROUP_FIELDS


def schema(document_type: str) -> dict[str, Any]:
    return {
        "scalars": SCALARS_BY_TYPE.get(document_type, PROPOSAL_SCALARS),
        "lists": {
            name: {"label": spec["label"], "fields": {f"{name}.{f}": label for f, label in spec["fields"].items()}}
            for name, spec in GROUPS.items()
            if not (name == "installment" and document_type != "POLICY")
        },
    }


# ------------------------------------------------------------------ leitura
def _open(data: bytes):
    from docx import Document

    try:
        return Document(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001 - qualquer ficheiro que o python-docx não abra
        raise InvalidTemplateError("O ficheiro não é um documento Word (.docx) válido") from exc


def _paragraph_text(paragraph) -> str:
    return "".join(run.text for run in paragraph.runs) if paragraph.runs else paragraph.text


def _iter_tables(container) -> Iterable:
    for table in container.tables:
        yield table
        for row in table.rows:
            for cell in row.cells:
                yield from _iter_tables(cell)


def _iter_paragraphs(container) -> Iterable:
    yield from container.paragraphs
    for table in container.tables:
        for row in table.rows:
            seen = set()
            for cell in row.cells:
                if id(cell._tc) in seen:
                    continue
                seen.add(id(cell._tc))
                yield from _iter_paragraphs(cell)


def _containers(document) -> list:
    parts = [document]
    for section in document.sections:
        for part in (section.header, section.footer, section.first_page_header, section.first_page_footer,
                     section.even_page_header, section.even_page_footer):
            parts.append(part)
    return parts


def placeholders_in(data: bytes) -> set[str]:
    document = _open(data)
    found: set[str] = set()
    for container in _containers(document):
        for paragraph in _iter_paragraphs(container):
            found.update(PLACEHOLDER.findall(_paragraph_text(paragraph)))
    return found


def validate(data: bytes, document_type: str) -> dict[str, Any]:
    """Recusa marcas desconhecidas (nunca sai um documento com `{{buraco}}`)."""
    found = placeholders_in(data)
    unknown = sorted(found - allowed_names(document_type))
    if unknown:
        raise InvalidTemplateError(
            f"O modelo Word usa marcas desconhecidas: {unknown}. "
            f"Consulte as marcas permitidas em GET /document-templates/schema."
        )
    if not found:
        raise InvalidTemplateError("O modelo Word não tem nenhuma marca {{variável}}: nada seria preenchido")
    return {"placeholders": sorted(found)}


# ----------------------------------------------------------------- preencher
def _replace_in_paragraph(paragraph, resolve) -> None:
    full = _paragraph_text(paragraph)
    if "{{" not in full:
        return
    runs = paragraph.runs
    if runs and sum(len(PLACEHOLDER.findall(r.text)) for r in runs) == len(PLACEHOLDER.findall(full)):
        # Cada marca está inteira num só bloco de texto: preserva o formato de cada um.
        for run in runs:
            if "{{" in run.text:
                run.text = PLACEHOLDER.sub(lambda m: resolve(m.group(1), m.group(0)), run.text)
        return
    new = PLACEHOLDER.sub(lambda m: resolve(m.group(1), m.group(0)), full)
    if runs:
        runs[0].text = new
        for run in runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(new)


def _group_of(text: str) -> str | None:
    for name in PLACEHOLDER.findall(text):
        head = name.split(".", 1)[0]
        if head in GROUPS and "." in name:
            return head
    return None


def _text_of_row(row) -> str:
    return " ".join(_paragraph_text(p) for cell in row.cells for p in cell.paragraphs)


def _group_resolver(group: str, item: dict[str, str]):
    def resolve(name: str, original: str) -> str:
        head, _, field = name.partition(".")
        if head == group:
            value = item.get(field)
            return "—" if value in (None, "") else str(value)
        return original  # outras marcas ficam para a passagem dos valores simples

    return resolve


def _expand_lists(document, groups: dict[str, list[dict[str, str]]]) -> None:
    from docx.oxml.ns import qn
    from docx.text.paragraph import Paragraph

    # Linhas de tabela
    for table in list(_iter_tables(document)):
        for row in list(table.rows):
            group = _group_of(_text_of_row(row))
            if group is None:
                continue
            tr = row._tr
            parent = tr.getparent()
            position = parent.index(tr)
            for offset, item in enumerate(groups.get(group, [])):
                clone = copy.deepcopy(tr)
                parent.insert(position + 1 + offset, clone)
                resolve = _group_resolver(group, item)
                for p in clone.iter(qn("w:p")):
                    _replace_in_paragraph(Paragraph(p, None), resolve)
            parent.remove(tr)
    # Parágrafos fora de tabelas
    for paragraph in list(document.paragraphs):
        group = _group_of(_paragraph_text(paragraph))
        if group is None:
            continue
        element = paragraph._p
        parent = element.getparent()
        position = parent.index(element)
        for offset, item in enumerate(groups.get(group, [])):
            clone = copy.deepcopy(element)
            parent.insert(position + 1 + offset, clone)
            _replace_in_paragraph(Paragraph(clone, None), _group_resolver(group, item))
        parent.remove(element)


def fill(data: bytes, *, scalars: dict[str, Any], groups: dict[str, list[dict[str, Any]]]) -> bytes:
    """Devolve o `.docx` com as marcas substituídas."""
    document = _open(data)
    clean_groups = {g: [{k: ("" if v is None else str(v)) for k, v in item.items()} for item in items] for g, items in groups.items()}
    _expand_lists(document, clean_groups)

    def resolve(name: str, original: str) -> str:
        if name not in scalars:
            return original
        value = scalars[name]
        return "—" if value in (None, "") else str(value)

    for container in _containers(document):
        for paragraph in _iter_paragraphs(container):
            _replace_in_paragraph(paragraph, resolve)
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


# --------------------------------------------------------------- dados
def groups_from_content(content: dict[str, Any], *, policy_members: dict[str, str] | None = None) -> dict[str, list[dict[str, Any]]]:
    """Listas a partir do conteúdo estruturado da proposta/apólice."""
    insured = []
    for person in content.get("population", []):
        status = person.get("status")
        if policy_members is not None:
            status = policy_members.get(person.get("memberId"), status)
        insured.append(
            {"name": person.get("name"), "relationship": person.get("relationship"),
             "age": person.get("age"), "status": status}
        )
    coverage = [
        {
            "name": c.get("name"),
            "limits": "; ".join(c.get("limits") or []),
            "participation": "; ".join(c.get("participation") or []),
            "waiting": c.get("waiting"),
            "priorApproval": "Sim" if c.get("requiresPriorApproval") else "Não",
        }
        for c in content.get("coverages", [])
        if c.get("contracted")
    ]
    return {
        "insured": insured,
        "coverage": coverage,
        "exclusion": [{"text": t} for t in content.get("exclusions", [])],
        "condition": [{"text": t} for t in content.get("conditions", [])],
        "waiting": [{"coverage": w.get("coverage"), "text": w.get("text")} for w in content.get("waiting", [])],
        "installment": [],
    }
