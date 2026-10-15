"""Teste de referência: o backend reproduz o cenário do separador
`Calculador` do Modelo_Actuarial_Saude_SAU_SAG_Angola.xlsx.

Cenário do Excel (células amarelas tal como vêm no ficheiro):
  SAU, plano Conforto (IND-MEDIO), Luanda, co-pagamento 5 000 Kz,
  comparticipação 80 %, franquias 0, sem venda cruzada, sem renovação;
  Titular M 38, Cônjuge F 35, Filho M 8, Filha F 5;
  opcionais escolhidos: Maternidade, Medicamentos, Estomatologia.

Valores esperados lidos das células calculadas do próprio Excel (não do
código). Diferenças só por arredondamento: o backend arredonda cada
componente a 2 casas (OD-006) e o Excel não arredonda nada.
"""
from __future__ import annotations

import datetime
import json
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import load_workbook

from apps.business_config.seed.actuarial_model import WORKBOOK_PATH, build_actuarial_model
from apps.foundation.application.context import CurrentActorProvider
from apps.foundation.domain.actor import Actor
from apps.foundation.domain.tenant import TenantId
from apps.policy.application.issuance_service import installment_amount
from apps.quote.infrastructure.models import Quote
from tests.quote.conftest import make_client, pin_pricing_day
from tests.quote.test_business_config_e2e import BASE, SCOPE, _publish, _values

EFFECTIVE_DATE = "2026-06-01"
PEOPLE = [  # (relação, sexo, data de nascimento -> idade em 2026-06-01, célula K do Excel)
    ("TITULAR", "M", "1988-01-15", "K25"),  # 38
    ("CONJUGE", "F", "1991-01-15", "K26"),  # 35
    ("FILHO", "M", "2018-01-15", "K27"),    # 8
    ("FILHO", "F", "2021-01-15", "K28"),    # 5
]


@pytest.fixture(scope="module")
def excel():
    sheet = load_workbook(Path(WORKBOOK_PATH), data_only=True)["Calculador"]
    cells = ("K25", "K26", "K27", "K28", "C64", "C70", "C84", "C85", "C86", "E91", "E92", "E93", "E94")
    return {ref: Decimal(str(sheet[ref].value)) for ref in cells}


@pytest.fixture
def approver(tenant):
    return make_client(tenant_id=tenant.id, roles=["ADMIN"], scope=SCOPE, actor_id="approver-1")


@pytest.mark.django_db
def test_backend_reproduces_the_excel_calculador(admin_client, approver, excel, tenant, monkeypatch):
    # Pressupostos do cenário que o teste fixa (se o Excel mudar, o teste avisa).
    sheet = load_workbook(Path(WORKBOOK_PATH), data_only=True)["Calculador"]
    assert (sheet["C5"].value, sheet["C6"].value, sheet["C8"].value) == ("SAU", "Conforto", "Luanda (Grande Luanda)")
    assert (sheet["C9"].value, sheet["C10"].value, sheet["C11"].value, sheet["C12"].value) == (5000, 0.8, 0, 0)
    assert (sheet["C13"].value, sheet["C14"].value) == ("Nenhuma", "Não")

    market = admin_client.post("/api/v1/product-templates/market", {}, format="json")
    assert market.status_code == 200, market.data
    plan = {p["code"]: p for p in admin_client.get("/api/v1/plans?size=50").data["items"]}["IND-MEDIO"]
    # O modelo de mercado (PDF comercial) não inclui doenças endémicas nos planos
    # individuais, mas a Matriz_SAU do Excel tem MAL = B no Conforto. Aqui o plano
    # é configurado como no Excel; a divergência está em docs/DECISOES_PENDENTES.md.
    endemic = next(i for i in admin_client.get("/api/v1/benefits?q=ENDEMIC_DISEASES&status=ALL&size=200").data["items"]
                   if i["code"] == "ENDEMIC_DISEASES")
    included = admin_client.put(f"/api/v1/plans/{plan['id']}/benefits/{endemic['id']}",
                                {"configuration": {"inclusion": "INCLUDED"}}, format="json")
    assert included.status_code in (200, 201), included.data
    for key, value in _values("1").items():
        _publish(admin_client, approver, key, value)
    _publish(admin_client, approver, "pricing.actuarial_model", build_actuarial_model())
    rule = admin_client.post(
        "/api/v1/rule-versions",
        {"code": "ELIG-XLS", "name": "Eligibility", "ruleDomain": "ELIGIBILITY", "effectiveFrom": "2026-01-01",
         "artifactType": "JSON", "artifactRef": json.dumps({})},
        format="json",
    ).data
    admin_client.post(f"/api/v1/configuration/rule-version/{rule['id']}/publish")

    proposer = admin_client.post(
        "/api/v1/proposers",
        {"proposerType": "INDIVIDUAL", "displayName": "Agregado Calculador", "registrationNumber": "123456789LA042",
         "dateOfBirth": "1988-01-15", "email": "calc@example.com", "phone": "+244900000000",
         "attributes": {"actuarialProfile": {
             "region": "Luanda (Grande Luanda)",
             "costSharing": {"copaymentAoa": 5000, "reimbursementShare": "0.8",
                             "outpatientDeductibleAoa": 0, "inpatientDeductibleAoa": 0},
             "crossSell": "Nenhuma",
         }}},
        format="json",
    )
    assert proposer.status_code == 201, proposer.data
    quote = admin_client.post(
        "/api/v1/quotes", {"proposerId": proposer.data["id"], "planId": plan["id"]},
        format="json",
    ).data
    # Idades do Excel calculadas em EFFECTIVE_DATE: fixa-se o dia do cálculo da cotação.
    pin_pricing_day(monkeypatch, datetime.date.fromisoformat(EFFECTIVE_DATE))
    member_cells = {}
    for relationship, sex, birth, cell in PEOPLE:
        member = admin_client.post(
            "/api/v1/members",
            {"proposerId": proposer.data["id"], "relationshipCode": relationship, "sexCode": sex, "dateOfBirth": birth},
            format="json",
        )
        assert member.status_code == 201, member.data
        member_cells[member.data["id"]] = cell
        admin_client.post(f"/api/v1/quotes/{quote['id']}/members", {"memberId": member.data["id"]}, format="json")
    current = admin_client.get(f"/api/v1/quotes/{quote['id']}").data
    selected = admin_client.put(
        f"/api/v1/quotes/{quote['id']}/selection",
        {"selection": {"optionalBenefitCodes": ["MATERNITY", "PHARMACY", "DENTAL"]}, "rowVersion": current["rowVersion"]},
        format="json",
    )
    assert selected.status_code == 200, selected.data

    priced = admin_client.post(f"/api/v1/quotes/{quote['id']}/price", {"recalculate": True}, format="json")
    assert priced.status_code == 200, priced.data
    components = priced.data["components"]

    # 1. Prémio puro base por pessoa (Calculador!K25:K28)
    for member_id, cell in member_cells.items():
        base = sum(Decimal(c["baseAmount"])
                   for c in components if c["memberId"] == member_id and c["code"].startswith("ACTUARIAL_")
                   and c["type"] == "BASE")
        assert abs(base - excel[cell]) <= Decimal("0.10"), (cell, base, excel[cell])

    # 2. Factor de risco combinado (C64) e desconto comercial limitado (C70)
    risk = {c["factorValue"] for c in components if c["code"] == "ACTUARIAL_RISK_FACTOR"}
    assert len(risk) == 1 and abs(Decimal(risk.pop()) - excel["C64"]) < Decimal("0.0000001")
    discount = {c["factorValue"] for c in components if c["code"] == "ACTUARIAL_COMMERCIAL_DISCOUNT"}
    assert len(discount) == 1 and Decimal(discount.pop()) - 1 == excel["C70"]

    # 3. Impostos (C85) e prémio total anual (C86)
    taxes = sum(Decimal(c["adjustmentAmount"]) for c in components if c["code"] in {"STAMP_DUTY", "ARSEG_LEVY"})
    assert abs(taxes - excel["C85"]) <= Decimal("0.05"), (taxes, excel["C85"])
    total = Decimal(priced.data["totalAmount"])
    assert abs(total - excel["C86"]) <= Decimal("0.10"), (total, excel["C86"])  # medido: +0,02 Kz

    # 4. Prestações com encargo de fracionamento (Calculador!E91:E94)
    token = CurrentActorProvider.bind(Actor(actor_id="test", tenant_id=TenantId(tenant.id), roles=frozenset()))
    try:
        quote_row = Quote.objects.select_related("latest_pricing_run").get(id=quote["id"])
        for frequency, cell in (("ANUAL", "E91"), ("SEMESTRAL", "E92"), ("TRIMESTRAL", "E93"), ("MENSAL", "E94")):
            amount = installment_amount(quote_row, total, frequency)
            assert abs(amount - excel[cell]) <= Decimal("0.10"), (frequency, amount, excel[cell])
    finally:
        CurrentActorProvider.reset(token)
