"""Modelos de documento com ficheiros: papel timbrado em PDF (A) e modelo Word com
marcas {{variável}} (B)."""
from __future__ import annotations

import datetime
import io

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile

from tests.quote.test_proposals import _generate
from tests.quote.test_quote_simulator import _scope

TODAY = datetime.date.today().isoformat()
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _letterhead_pdf(text="PAPEL TIMBRADO SEGURADORA ANGOLA") -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    buffer = io.BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=A4)
    pdf.setFont("Helvetica-Bold", 14)
    pdf.drawString(60, A4[1] - 50, text)
    pdf.setFont("Helvetica", 9)
    pdf.drawString(60, 40, "Rodape da seguradora - Luanda")
    pdf.save()
    return buffer.getvalue()


def _word(paragraphs=(), table_row=None, header_cell=None) -> bytes:
    from docx import Document

    document = Document()
    for text in paragraphs:
        document.add_paragraph(text)
    if table_row:
        table = document.add_table(rows=2, cols=len(table_row))
        for index, label in enumerate(("Pessoa", "Relação")[: len(table_row)]):
            table.rows[0].cells[index].text = label
        for index, value in enumerate(table_row):
            table.rows[1].cells[index].text = value
    if header_cell:
        document.sections[0].header.paragraphs[0].text = header_cell
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def _new_template(client, **extra):
    response = client.post(
        "/api/v1/document-templates", {"effectiveFrom": TODAY, **extra}, format="json"
    )
    assert response.status_code == 201, response.data
    return response.data


def _text_of_pdf(data: bytes) -> list[str]:
    from pypdf import PdfReader

    return [page.extract_text() for page in PdfReader(io.BytesIO(data)).pages]


@pytest.mark.django_db
def test_template_can_be_created_without_json_and_gets_defaults(admin_client):
    created = _new_template(admin_client)
    assert created["templateCode"] == "QUOTE_PROPOSAL_DEFAULT" and created["documentType"] == "QUOTE_PROPOSAL"
    assert created["hasLetterhead"] is False and created["hasWordTemplate"] is False
    policy = _new_template(admin_client, documentType="POLICY")
    assert policy["templateCode"] == "POLICY_DEFAULT" and policy["documentType"] == "POLICY"
    schema = admin_client.get("/api/v1/document-templates/schema").data["word"]
    assert "client.name" in schema["QUOTE_PROPOSAL"]["scalars"] and "policy.number" in schema["POLICY"]["scalars"]
    assert "insured" in schema["QUOTE_PROPOSAL"]["lists"] and "installment" in schema["POLICY"]["lists"]
    assert "installment" not in schema["QUOTE_PROPOSAL"]["lists"]


@pytest.mark.django_db
def test_letterhead_pdf_becomes_the_background_of_the_proposal(admin_client):
    scope = _scope(admin_client)
    template = _new_template(admin_client)
    uploaded = admin_client.post(
        f"/api/v1/document-templates/{template['id']}/letterhead",
        {"file": SimpleUploadedFile("timbrado.pdf", _letterhead_pdf(), content_type="application/pdf"), "top": "3", "bottom": "2.5"},
        format="multipart",
    )
    assert uploaded.status_code == 200, uploaded.data
    assert uploaded.data["hasLetterhead"] is True
    assert uploaded.data["letterhead"]["marginsCm"]["top"] == 3.0 and uploaded.data["letterhead"]["marginsCm"]["left"] == 2.0
    assert admin_client.get(f"/api/v1/document-templates/{template['id']}/letterhead").content.startswith(b"%PDF")

    # A pré-visualização do modelo já mostra o papel timbrado
    preview = admin_client.get(f"/api/v1/document-templates/{template['id']}/preview", {"fileType": "pdf"})
    assert preview.status_code == 200 and "PAPEL TIMBRADO SEGURADORA ANGOLA" in _text_of_pdf(preview.content)[0]

    admin_client.post(f"/api/v1/document-templates/{template['id']}/publish")
    generated = _generate(admin_client, scope).data
    assert generated["status"] == "GENERATED", generated
    pages = _text_of_pdf(admin_client.get(generated["downloadUrl"]).content)
    assert "PAPEL TIMBRADO SEGURADORA ANGOLA" in pages[0] and "Rodape da seguradora" in pages[0]
    assert "Proposta" in pages[0]  # o conteúdo da proposta está por cima do papel
    assert all("PAPEL TIMBRADO" in page for page in pages)  # o fundo repete-se em todas as páginas

    # publicado = imutável
    locked = admin_client.post(
        f"/api/v1/document-templates/{template['id']}/letterhead",
        {"file": SimpleUploadedFile("outro.pdf", _letterhead_pdf("OUTRO"), content_type="application/pdf")}, format="multipart",
    )
    assert locked.status_code == 409 and "rascunho" in locked.data["detail"]


@pytest.mark.django_db
def test_word_template_variables_and_repeated_rows_are_filled(admin_client):
    scope = _scope(admin_client)
    template = _new_template(admin_client)
    data = _word(
        paragraphs=["Proposta para {{client.name}} (NIF {{client.nif}}) - plano {{plan.name}}", "Exclusão: {{exclusion.text}}"],
        table_row=["{{insured.name}}", "{{insured.relationship}}"],
        header_cell="Seguradora Exemplo - {{document.number}}",
    )
    uploaded = admin_client.post(
        f"/api/v1/document-templates/{template['id']}/word",
        {"file": SimpleUploadedFile("modelo.docx", data, content_type=DOCX)}, format="multipart",
    )
    assert uploaded.status_code == 200, uploaded.data
    assert set(uploaded.data["word"]["placeholders"]) >= {"client.name", "insured.name", "exclusion.text", "document.number"}
    admin_client.post(f"/api/v1/document-templates/{template['id']}/publish")

    generated = _generate(admin_client, scope, format="DOCX").data
    assert generated["status"] == "GENERATED", generated
    from docx import Document

    document = Document(io.BytesIO(admin_client.get(generated["downloadUrl"]).content))
    texts = [p.text for p in document.paragraphs]
    assert any(t.startswith("Proposta para Holder (NIF") and "{{" not in t for t in texts)
    exclusions = [t for t in texts if t.startswith("Exclusão: ")]
    assert len(exclusions) >= 5 and all("{{" not in t for t in exclusions)  # um parágrafo por exclusão
    table = document.tables[0]
    assert len(table.rows) == 3  # cabeçalho + uma linha por pessoa segura (2)
    assert all("{{" not in cell.text for row in table.rows for cell in row.cells)
    assert generated["documentNumber"] in document.sections[0].header.paragraphs[0].text


@pytest.mark.django_db
def test_template_file_validation_is_in_portuguese(admin_client):
    template = _new_template(admin_client)
    base = f"/api/v1/document-templates/{template['id']}"

    unknown = admin_client.post(
        f"{base}/word", {"file": SimpleUploadedFile("m.docx", _word(["{{cliente.nome}}"]), content_type=DOCX)}, format="multipart"
    )
    assert unknown.status_code == 400 and "marcas desconhecidas" in unknown.data["detail"] and "cliente.nome" in unknown.data["detail"]

    none = admin_client.post(
        f"{base}/word", {"file": SimpleUploadedFile("m.docx", _word(["sem marcas"]), content_type=DOCX)}, format="multipart"
    )
    assert none.status_code == 400 and "nenhuma marca" in none.data["detail"]

    not_docx = admin_client.post(
        f"{base}/word", {"file": SimpleUploadedFile("m.docx", b"isto nao e word", content_type=DOCX)}, format="multipart"
    )
    assert not_docx.status_code == 400 and "Word" in not_docx.data["detail"]

    wrong_type = admin_client.post(
        f"{base}/letterhead", {"file": SimpleUploadedFile("m.txt", b"texto", content_type="text/plain")}, format="multipart"
    )
    assert wrong_type.status_code == 400 and "PDF" in wrong_type.data["detail"]

    broken = admin_client.post(
        f"{base}/letterhead", {"file": SimpleUploadedFile("m.pdf", b"%PDF-1.4 lixo", content_type="application/pdf")}, format="multipart"
    )
    assert broken.status_code == 400

    bad_margin = admin_client.post(
        f"{base}/letterhead",
        {"file": SimpleUploadedFile("t.pdf", _letterhead_pdf(), content_type="application/pdf"), "top": "99"}, format="multipart",
    )
    assert bad_margin.status_code == 400 and "margem" in bad_margin.data["detail"]

    # remover
    admin_client.post(f"{base}/letterhead", {"file": SimpleUploadedFile("t.pdf", _letterhead_pdf(), content_type="application/pdf")}, format="multipart")
    removed = admin_client.delete(f"{base}/letterhead")
    assert removed.status_code == 200 and removed.data["hasLetterhead"] is False
