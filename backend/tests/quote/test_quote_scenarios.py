"""R5 — criar, clonar, comparar e selecionar cenários de cotação; aplicar um
cenário à cotação. O motor é o mesmo da cotação (pré-visualização): nada é
gravado no motor e a cotação oficial só muda com `apply`."""
from __future__ import annotations

import io
from decimal import Decimal

import pytest

from tests.quote.conftest import make_client
from tests.quote.test_quote_simulator import _quote, _runs, _scope, _total


def _create(client, scope, name, parameters=None, description=None):
    body = {"name": name}
    if parameters is not None:
        body["parameters"] = parameters
    if description:
        body["description"] = description
    return client.post(f"/api/v1/quotes/{scope['quoteId']}/scenarios", body, format="json")


@pytest.mark.django_db
def test_new_scenario_starts_from_the_quote_state_and_is_calculated(admin_client):
    scope = _scope(admin_client)
    runs = _runs(admin_client)
    created = _create(admin_client, scope, "Base")
    assert created.status_code == 201, created.data
    # a opção por omissão do plano fica explícita no cenário (congelada)
    assert created.data["parameters"] == {"selection": {"participationOption": "B"}, "adjustments": []}
    assert created.data["result"]["totalAmount"] == "200.00"
    assert created.data["result"]["currencyCode"] == "AOA"
    assert created.data["result"]["memberCount"] == 2
    assert created.data["isSelected"] is False and created.data["stale"] is False
    assert _runs(admin_client) == runs  # o motor não grava nada
    assert _total(admin_client, scope) == Decimal("200.00")  # a cotação oficial não mudou


@pytest.mark.django_db
def test_scenarios_with_different_parameters_have_different_premiums(admin_client):
    scope = _scope(admin_client)
    optimistic = _create(admin_client, scope, "Participação elevada", {"selection": {"participationOption": "D"}})
    pessimistic = _create(
        admin_client, scope, "Sem participação + dentária",
        {"selection": {"participationOption": "A", "optionalBenefitCodes": ["DENTAL"]},
         "adjustments": [{"kind": "SURCHARGE_PERCENT", "value": "10", "reason": "Sinistralidade elevada"}]},
    )
    assert optimistic.data["result"]["totalAmount"] == "176.00"  # 88 * 2
    # A (+15%) = 115, + 10 dentária = 125, +10% = 137.50 por pessoa
    assert pessimistic.data["result"]["totalAmount"] == "275.00"
    assert pessimistic.data["parameters"]["adjustments"][0]["value"] == "10"
    assert _quote(admin_client, scope)["selection"] == {}


@pytest.mark.django_db
def test_scenario_name_must_be_unique_per_quote_and_parameters_are_validated(admin_client):
    scope = _scope(admin_client)
    assert _create(admin_client, scope, "A").status_code == 201
    dup = _create(admin_client, scope, "A")
    assert dup.status_code == 409 and dup.data["code"] == "CONFLICT"
    for bad in (
        {"selection": {"participationOption": "Z"}},
        {"selection": {"optionalBenefitCodes": ["OPTICAL"]}},
        {"adjustments": [{"kind": "SURCHARGE_PERCENT", "value": "0"}]},
        {"adjustments": [{"kind": "SURCHARGE_AMOUNT", "value": "5", "scope": "MEMBER"}]},
        {"unknown": 1},
    ):
        assert _create(admin_client, scope, "B", bad).status_code == 400, bad


@pytest.mark.django_db
def test_clone_copies_parameters_and_links_to_the_original(admin_client):
    scope = _scope(admin_client)
    original = _create(admin_client, scope, "Cenário base", {"selection": {"participationOption": "C"}}).data
    clone = admin_client.post(f"/api/v1/scenarios/{original['id']}/clone", {}, format="json")
    assert clone.status_code == 201
    assert clone.data["name"] == "Cenário base (cópia)"
    assert clone.data["baseScenarioId"] == original["id"]
    assert clone.data["parameters"] == original["parameters"]
    assert clone.data["result"]["totalAmount"] == original["result"]["totalAmount"]
    second = admin_client.post(f"/api/v1/scenarios/{original['id']}/clone", {}, format="json")
    assert second.data["name"] == "Cenário base (cópia) 2"
    named = admin_client.post(f"/api/v1/scenarios/{original['id']}/clone", {"name": "Pessimista"}, format="json")
    assert named.data["name"] == "Pessimista"
    assert admin_client.post(f"/api/v1/scenarios/{original['id']}/clone", {"name": "Pessimista"}, format="json").status_code == 409


@pytest.mark.django_db
def test_update_scenario_recalculates_and_uses_optimistic_locking(admin_client):
    scope = _scope(admin_client)
    scenario = _create(admin_client, scope, "Editável").data
    updated = admin_client.patch(
        f"/api/v1/scenarios/{scenario['id']}",
        {"name": "Editável v2", "parameters": {"selection": {"participationOption": "A"}}, "rowVersion": scenario["rowVersion"]},
        format="json",
    )
    assert updated.status_code == 200, updated.data
    assert updated.data["name"] == "Editável v2" and updated.data["result"]["totalAmount"] == "230.00"
    stale = admin_client.patch(
        f"/api/v1/scenarios/{scenario['id']}", {"name": "x", "rowVersion": scenario["rowVersion"]}, format="json"
    )
    assert stale.status_code == 409 and stale.data["code"] == "VERSION-CONFLICT"
    # sem `parameters`: mantém os parâmetros, só renomeia
    renamed = admin_client.patch(
        f"/api/v1/scenarios/{scenario['id']}", {"name": "Só nome", "rowVersion": updated.data["rowVersion"]}, format="json"
    )
    assert renamed.data["parameters"]["selection"] == {"participationOption": "A"}


@pytest.mark.django_db
def test_compare_scenarios_side_by_side(admin_client):
    scope = _scope(admin_client)
    base = _create(admin_client, scope, "Base").data
    a = _create(admin_client, scope, "Opção A + dentária", {"selection": {"participationOption": "A", "optionalBenefitCodes": ["DENTAL"]}}).data
    hospital = {"coverageOverrides": {"HOSPITALIZATION": {"limits": [{"kind": "ANNUAL_PER_PERSON", "amount": "40000000"}]}}}
    big = _create(admin_client, scope, "Capital maior", {"selection": hospital}).data

    ids = ",".join([base["id"], a["id"], big["id"]])
    result = admin_client.get(f"/api/v1/quotes/{scope['quoteId']}/scenarios/compare?ids={ids}")
    assert result.status_code == 200, result.data
    data = result.data
    assert data["currencyCode"] == "AOA"
    totals = {s["name"]: s["totalAmount"] for s in data["scenarios"]}
    assert totals == {"Base": "200.00", "Opção A + dentária": "250.00", "Capital maior": "200.00"}
    by_name = {s["name"]: s for s in data["scenarios"]}
    assert by_name["Base"]["deltaVsFirst"] == "0.00"
    assert by_name["Opção A + dentária"]["deltaVsFirst"] == "50.00" and by_name["Opção A + dentária"]["deltaPercentVsFirst"] == "25.00"
    assert by_name["Opção A + dentária"]["participationOption"] == "A"
    assert by_name["Opção A + dentária"]["optionalCoverages"] == ["DENTAL"]
    assert data["summary"]["lowest"]["totalAmount"] == "200.00"
    assert data["summary"]["highest"]["name"] == "Opção A + dentária"
    assert data["summary"]["spread"] == "50.00"

    coverages = {c["code"]: c for c in data["coverages"]}
    # só a cobertura dentária (opcional contratada) e a hospitalização (capital) diferem
    assert coverages["DENTAL"]["differs"] is True
    assert coverages["DENTAL"]["values"][a["id"]]["contracted"] is True
    assert coverages["DENTAL"]["values"][base["id"]]["contracted"] is False
    assert coverages["HOSPITALIZATION"]["differs"] is True
    assert coverages["HOSPITALIZATION"]["values"][big["id"]]["limits"][0]["amount"] == "40000000"
    assert coverages["EMERGENCY"]["differs"] is False

    components = {c["code"]: c for c in data["components"]}
    assert components["PARTICIPATION_OPTION_A"]["differs"] is True
    assert components["COVERAGE_DENTAL"]["values"][a["id"]]["amount"] == "20.00"
    assert components["COVERAGE_DENTAL"]["values"][base["id"]] is None


@pytest.mark.django_db
def test_compare_needs_two_scenarios_of_the_same_quote(admin_client):
    scope = _scope(admin_client)
    only = _create(admin_client, scope, "Único").data
    url = f"/api/v1/quotes/{scope['quoteId']}/scenarios/compare"
    assert admin_client.get(url).status_code == 400
    assert admin_client.get(f"{url}?ids={only['id']}").status_code == 400
    import uuid

    assert admin_client.get(f"{url}?ids={only['id']},{uuid.uuid4()}").status_code == 404
    assert admin_client.get(f"{url}?ids=not-a-uuid").status_code == 403


@pytest.mark.django_db
def test_compare_export_to_xlsx_and_csv(admin_client):
    scope = _scope(admin_client)
    _create(admin_client, scope, "Base")
    _create(admin_client, scope, "Opção D", {"selection": {"participationOption": "D"}})
    url = f"/api/v1/quotes/{scope['quoteId']}/scenarios/compare/export"
    xlsx = admin_client.get(f"{url}?fileType=xlsx")
    assert xlsx.status_code == 200
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(xlsx.content))
    assert workbook.sheetnames == ["Resumo", "Coberturas", "Componentes"]
    summary = [[c.value for c in row] for row in workbook["Resumo"].iter_rows()]
    assert summary[0][:3] == ["Cenário", "Prémio total", "Δ vs 1.º"]
    assert summary[1][0] == "Base" and summary[2][0] == "Opção D" and summary[2][1] == "176.00"
    csv_resp = admin_client.get(f"{url}?fileType=csv")
    assert csv_resp.status_code == 200 and "# Resumo" in csv_resp.content.decode("utf-8-sig")
    assert admin_client.get(f"{url}?fileType=pdf").status_code == 400


@pytest.mark.django_db
def test_select_one_scenario_per_quote_and_cannot_delete_the_selected_one(admin_client):
    scope = _scope(admin_client)
    a = _create(admin_client, scope, "A").data
    b = _create(admin_client, scope, "B").data
    assert admin_client.post(f"/api/v1/scenarios/{a['id']}/select", {}, format="json").data["isSelected"] is True
    assert admin_client.post(f"/api/v1/scenarios/{b['id']}/select", {}, format="json").data["isSelected"] is True
    listed = {s["name"]: s["isSelected"] for s in admin_client.get(f"/api/v1/quotes/{scope['quoteId']}/scenarios").data}
    assert listed == {"A": False, "B": True}
    assert admin_client.delete(f"/api/v1/scenarios/{b['id']}").status_code == 409
    assert admin_client.delete(f"/api/v1/scenarios/{a['id']}").status_code == 204
    assert admin_client.get(f"/api/v1/scenarios/{a['id']}").status_code == 404


@pytest.mark.django_db
def test_scenario_goes_stale_when_the_plan_or_population_changes(admin_client):
    scope = _scope(admin_client)
    scenario = _create(admin_client, scope, "Base").data
    assert scenario["stale"] is False

    # o GTE altera o capital da hospitalização no plano
    benefit = next(i for i in admin_client.get("/api/v1/benefits?q=HOSPITALIZATION&size=200").data["items"] if i["code"] == "HOSPITALIZATION")
    admin_client.put(
        f"/api/v1/plans/{scope['planId']}/benefits/{benefit['id']}",
        {"configuration": {"limits": [{"kind": "ANNUAL_PER_PERSON", "amount": "20000000"}]}},
        format="json",
    )
    assert admin_client.get(f"/api/v1/scenarios/{scenario['id']}").data["stale"] is True
    recalculated = admin_client.post(f"/api/v1/scenarios/{scenario['id']}/recalculate", {}, format="json")
    assert recalculated.status_code == 200 and recalculated.data["stale"] is False
    compare_after = _create(admin_client, scope, "Outro")
    assert compare_after.data["stale"] is False


@pytest.mark.django_db
def test_apply_scenario_makes_it_the_official_state(admin_client):
    scope = _scope(admin_client)
    admin_client.post(
        f"/api/v1/quotes/{scope['quoteId']}/adjustments",
        {"kind": "DISCOUNT_PERCENT", "value": "5", "reason": "Desconto antigo da cotação"},
        format="json",
    )
    assert _total(admin_client, scope) == Decimal("190.00")
    scenario = _create(
        admin_client, scope, "Proposta final",
        {"selection": {"participationOption": "C", "optionalBenefitCodes": ["DENTAL"]},
         "adjustments": [{"kind": "SURCHARGE_PERCENT", "value": "10", "reason": "Agravamento do cenário final"}]},
    ).data
    # C: 93 + 10 = 103; +10% = 113.30 por pessoa
    assert scenario["result"]["totalAmount"] == "226.60"

    short = admin_client.post(f"/api/v1/scenarios/{scenario['id']}/apply", {"reason": "x"}, format="json")
    assert short.status_code == 400
    applied = admin_client.post(f"/api/v1/scenarios/{scenario['id']}/apply", {"reason": "Cliente escolheu esta opção"}, format="json")
    assert applied.status_code == 200, applied.data
    assert applied.data["status"] == "READY"
    assert applied.data["selection"] == {"optionalBenefitCodes": ["DENTAL"], "participationOption": "C"}
    assert _total(admin_client, scope) == Decimal("226.60")

    adjustments = admin_client.get(f"/api/v1/quotes/{scope['quoteId']}/adjustments").data
    assert [(a["kind"], a["status"]) for a in adjustments] == [("DISCOUNT_PERCENT", "REMOVED"), ("SURCHARGE_PERCENT", "ACTIVE")]
    assert "Proposta final" in adjustments[0]["removedReason"]
    assert admin_client.get(f"/api/v1/scenarios/{scenario['id']}").data["stale"] is False


@pytest.mark.django_db
def test_apply_scenario_respects_adjustment_authority_limits(admin_client, tenant):
    from apps.quote.infrastructure.models import AuthorityPolicy

    scope = _scope(admin_client)
    scenario = _create(
        admin_client, scope, "Agravamento alto",
        {"adjustments": [{"kind": "SURCHARGE_PERCENT", "value": "20", "reason": "Agravamento alto"}]},
    ).data
    AuthorityPolicy._base_manager.create(
        tenant_id=tenant.id, code="AUTH", version_no=1, effective_from="2026-01-01", status="PUBLISHED",
        policy_json={"maxAdjustmentPct": "5"},
    )
    subscriber = make_client(tenant_id=tenant.id, roles=["SUBSCRIBER"], scope="q360.read q360.write", actor_id="sub-1")
    denied = subscriber.post(f"/api/v1/scenarios/{scenario['id']}/apply", {"reason": "Tentativa sem autoridade"}, format="json")
    assert denied.status_code == 403
    assert _total(admin_client, scope) == Decimal("200.00")  # nada mudou
    ok = admin_client.post(f"/api/v1/scenarios/{scenario['id']}/apply", {"reason": "Aprovado pelo responsável"}, format="json")
    assert ok.status_code == 200 and _total(admin_client, scope) == Decimal("240.00")


@pytest.mark.django_db
def test_scenarios_are_frozen_once_the_quote_is_no_longer_editable(admin_client):
    scope = _scope(admin_client)
    scenario = _create(admin_client, scope, "Base").data
    quote = _quote(admin_client, scope)
    admin_client.post(f"/api/v1/quotes/{scope['quoteId']}/cancel", {"reasonCode": "TEST", "reasonText": "teste", "rowVersion": quote["rowVersion"]}, format="json")
    assert _create(admin_client, scope, "Novo").status_code == 409
    assert admin_client.post(f"/api/v1/scenarios/{scenario['id']}/clone", {}, format="json").status_code == 409
    assert admin_client.get(f"/api/v1/scenarios/{scenario['id']}").status_code == 200  # leitura continua


@pytest.mark.django_db
def test_scenario_changes_are_audited(admin_client, tenant):
    scope = _scope(admin_client)
    scenario = _create(admin_client, scope, "Auditado", {"selection": {"participationOption": "D"}}).data
    admin_client.post(f"/api/v1/scenarios/{scenario['id']}/select", {}, format="json")
    auditor = make_client(tenant_id=tenant.id, roles=[], scope="q360.read q360.audit")
    events = auditor.get(f"/api/v1/audit-events?entityId={scenario['id']}&size=100").data["items"]
    actions = {e["action"] for e in events}
    assert {"quote.scenario_created", "quote.scenario_selected"} <= actions
    created = next(e for e in events if e["action"] == "quote.scenario_created")
    assert created["metadata"]["after"]["totalAmount"] == "176.00"


@pytest.mark.django_db
def test_readers_can_see_but_not_change_scenarios(admin_client, tenant):
    scope = _scope(admin_client)
    scenario = _create(admin_client, scope, "Visível").data
    reader = make_client(tenant_id=tenant.id, roles=[], scope="q360.read q360.audit")
    assert reader.get(f"/api/v1/quotes/{scope['quoteId']}/scenarios").status_code == 200
    assert reader.post(f"/api/v1/quotes/{scope['quoteId']}/scenarios", {"name": "X"}, format="json").status_code == 403
    assert reader.post(f"/api/v1/scenarios/{scenario['id']}/select", {}, format="json").status_code == 403
