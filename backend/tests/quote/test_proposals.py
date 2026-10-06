"""R8 — propostas em PDF/DOCX/HTML a partir da cotação ou de um cenário, com
modelos de documento configuráveis pelo GTE."""
from __future__ import annotations

import hashlib
import io
import json
from pathlib import Path

import pytest

from tests.quote.conftest import make_client
from tests.quote.test_quote_simulator import SETTINGS, _benefit, _quote, _scope

REASON = "Motivo de teste suficiente"
INDICATIVE = {**SETTINGS, "indicative": True, "accessSystem": "MIXED", "entryAge": {"min": 0, "max": 65},
              "permanenceAge": {"max": 75}, "priorAuthorization": {"internments": True}, "network": {"regions": ["LUANDA"]},
              "preExistingTreatment": {"mode": "TOTAL_EXCLUSION"}, "household": {"spouse": True, "dependentChildMaxAge": 25}}


def _generate(client, scope, **body):
    return client.post(f"/api/v1/quotes/{scope['quoteId']}/proposals", {"format": "PDF", **body}, format="json")


def _html(client, scope, **query):
    qs = "&".join(f"{k}={v}" for k, v in query.items())
    resp = client.get(f"/api/v1/quotes/{scope['quoteId']}/proposals/preview?{qs}")
    assert resp.status_code == 200, resp.content
    return resp.content.decode("utf-8")


def _configure(client, scope):
    """Capitais e participação reais para a proposta mostrar."""
    hospital = _benefit(client, "HOSPITALIZATION")
    client.put(
        f"/api/v1/plans/{scope['planId']}/benefits/{hospital['id']}",
        {"configuration": {"limits": [{"kind": "ANNUAL_PER_PERSON", "amount": "15000000"},
                                      {"kind": "PER_ACT_OR_DAY", "label": "Quarto por dia", "amount": "50000"}],
                           "participation": {"NETWORK": {"type": "NONE", "perEventDeductible": "100000"},
                                             "REIMBURSEMENT": {"type": "COINSURANCE", "insurerPct": "85", "insuredAnnualCap": "500000"}},
                           "waitingDays": {"individual": 90, "scheduledSurgery": 180, "accidentWaived": True}}},
        format="json",
    )
    evacuation = _benefit(client, "MEDICAL_EVACUATION")
    client.put(
        f"/api/v1/plans/{scope['planId']}/benefits/{evacuation['id']}",
        {"configuration": {"inclusion": "OPTIONAL", "currency": "USD", "destinations": ["Portugal", "Brasil"],
                           "limits": [{"kind": "ANNUAL_PER_PERSON", "amount": "100000"}]}},
        format="json",
    )


@pytest.mark.django_db
def test_generate_a_pdf_proposal_with_the_default_template(admin_client):
    scope = _scope(admin_client)
    created = _generate(admin_client, scope)
    assert created.status_code == 201, created.data
    doc = created.data
    assert doc["status"] == "GENERATED" and doc["format"] == "PDF" and doc["proposalNo"] == 1
    assert doc["templateCode"] == "QUOTE_PROPOSAL_DEFAULT" and doc["templateVersion"] == 1
    assert doc["documentNumber"] and doc["fileName"].endswith("-1.pdf") and doc["contentType"] == "application/pdf"
    download = admin_client.get(doc["downloadUrl"])
    assert download.status_code == 200 and download.content.startswith(b"%PDF")
    assert hashlib.sha256(download.content).hexdigest() == doc["checksum"]
    assert "attachment" in download["Content-Disposition"] and doc["fileName"] in download["Content-Disposition"]
    # o modelo por omissão ficou publicado e editável como os outros
    templates = admin_client.get("/api/v1/document-templates?templateCode=QUOTE_PROPOSAL_DEFAULT").data["items"]
    assert [(t["templateCode"], t["status"]) for t in templates] == [("QUOTE_PROPOSAL_DEFAULT", "PUBLISHED")]
    # histórico
    second = _generate(admin_client, scope, format="docx")
    assert second.data["proposalNo"] == 2 and second.data["format"] == "DOCX"
    listed = admin_client.get(f"/api/v1/quotes/{scope['quoteId']}/proposals").data
    assert [d["proposalNo"] for d in listed] == [2, 1]
    assert doc["documentNumber"] != second.data["documentNumber"]


@pytest.mark.django_db
def test_docx_and_html_carry_the_same_content(admin_client):
    scope = _scope(admin_client, settings=INDICATIVE)
    _configure(admin_client, scope)
    docx = _generate(admin_client, scope, format="DOCX").data
    raw = admin_client.get(docx["downloadUrl"]).content
    assert raw.startswith(b"PK")
    from docx import Document

    document = Document(io.BytesIO(raw))
    text = "\n".join(p.text for p in document.paragraphs) + "\n".join(c.text for t in document.tables for r in t.rows for c in r.cells)
    for expected in ("Holder", "SIM", "Hospitalização", "Capital anual por pessoa: 15.000.000 Kz", "Evacuação e tratamento no estrangeiro"):
        assert expected in text
    html_doc = _generate(admin_client, scope, format="HTML").data
    html = admin_client.get(html_doc["downloadUrl"]).content.decode("utf-8")
    assert "Capital anual por pessoa: 15.000.000 Kz" in html and html_doc["contentType"].startswith("text/html")


@pytest.mark.django_db
def test_proposal_content_follows_the_pdf_rules(admin_client):
    scope = _scope(admin_client, settings=INDICATIVE)
    _configure(admin_client, scope)
    first = scope["memberIds"][0]
    admin_client.put(
        f"/api/v1/quotes/{scope['quoteId']}/selection",
        {"selection": {"participationOption": "C", "optionalBenefitCodes": ["DENTAL"]}, "rowVersion": _quote(admin_client, scope)["rowVersion"]},
        format="json",
    )
    admin_client.post(
        f"/api/v1/quotes/{scope['quoteId']}/adjustments",
        {"kind": "DISCOUNT_PERCENT", "value": "5", "reason": "Desconto comercial de fidelização"}, format="json",
    )
    admin_client.put(
        f"/api/v1/quotes/{scope['quoteId']}/members/{first}/pre-existing",
        {"decision": {"mode": "ACCEPTED_WITH_LOADING", "loadingPct": "25", "condition": "Hipertensão", "clinicalQuestionnaireCompleted": True},
         "reason": "Aceite após análise clínica"}, format="json",
    )
    html = _html(admin_client, scope)

    # coberturas: capitais, sublimite, participação rede/reembolso, carência, destinos, moeda
    assert "Capital anual por pessoa: 15.000.000 Kz" in html and "Quarto por dia: 50.000 Kz" in html
    assert "Rede convencionada: sem participação do segurado; franquia de 100.000 Kz por evento" in html
    assert "Reembolso: a seguradora paga 85% e o segurado 15% (teto anual a cargo do segurado: 500.000 Kz)" in html
    assert "Carência: 90 dias; cirurgias programadas: 180 dias; em acidente não há carência" in html
    assert "Dentária" not in html and "Estomatologia" in html  # opcional escolhida
    assert "Estomatologia e ortodontia (opcional escolhida)" in html
    assert "Coberturas opcionais disponíveis (não incluídas)" in html and "Evacuação e tratamento no estrangeiro" in html
    # opção de participação A–E
    assert "C — Participação moderada" in html and "Franquia anual" in html and "25.000 Kz" in html
    # preço e ajustes
    assert "Desconto de 5%" in html and "Prémio total" in html
    # preexistências por pessoa
    assert "Aceitação com sobreprémio" in html and "Hipertensão" in html and "Sobreprémio: 25%" in html
    assert "Tratamento previsto no plano: Exclusão total" in html
    # condições do PDF
    assert "Pré-autorização obrigatória para: internamentos" in html
    assert "Idade de entrada até 65 anos; renovação até 75 anos." in html
    assert "Agregado familiar: cônjuge ou unido de facto, filhos até aos 25 anos." in html
    assert "Rede convencionada em: Luanda." in html
    assert "Não há regra de rateio" in html
    # exclusões habituais (13) e aviso de valores indicativos
    for text in ("Cirurgia plástica ou estética", "Tratamentos de infertilidade", "Desportos de alto risco", "Tratamentos no estrangeiro sem aprovação prévia", "Internamentos em lares"):
        assert text in html
    assert "Os valores em Kz são indicativos" in html
    assert "Carência" in html and "Períodos de carência" in html and "Questionário clínico" in html


@pytest.mark.django_db
def test_group_proposal_shows_group_conditions(admin_client):
    group = {"segment": "GROUP", "tier": "CUSTOM", "category": "Quadros", "accessSystem": "MIXED", "indicative": True,
             "group": {"minInsured": 1, "adhesion": "MANDATORY", "premiumPayer": "EMPLOYEE_COPAYS_FAMILY", "monthlyInclusionsExclusions": True,
                       "profitSharing": {"lossRatioThresholdPct": "65"}, "continuityAfterLeaving": True, "includesRetirees": True},
             "preExistingTreatment": {"mode": "FULL_ACCEPTANCE"}}
    scope = _scope(admin_client, settings=group)
    html = _html(admin_client, scope)
    for text in ("Grupo (Empresas)", "Quadros", "Adesão obrigatória.", "copago pelo trabalhador", "Inclusões e exclusões mensais",
                 "Participação nos resultados", "65%", "Continuidade após saída da empresa", "Inclui reformados.",
                 "Aceitação integral"):
        assert text in html, text


@pytest.mark.django_db
def test_selected_scenario_drives_the_proposal(admin_client):
    scope = _scope(admin_client)
    base = admin_client.post(f"/api/v1/quotes/{scope['quoteId']}/scenarios", {"name": "Base"}, format="json").data
    big = admin_client.post(
        f"/api/v1/quotes/{scope['quoteId']}/scenarios",
        {"name": "Opção A", "parameters": {"selection": {"participationOption": "A"}}}, format="json",
    ).data
    assert "200 Kz" in _html(admin_client, scope)  # sem selecionado: estado atual da cotação
    admin_client.post(f"/api/v1/scenarios/{big['id']}/select", {}, format="json")
    html = _html(admin_client, scope)
    assert "230 Kz" in html and "Opção A" in html and "A — Sem participação" in html
    explicit = _generate(admin_client, scope, scenarioId=base["id"])
    assert explicit.data["scenarioId"] == base["id"]
    default = _generate(admin_client, scope, format="HTML")
    assert default.data["scenarioId"] == big["id"]


@pytest.mark.django_db
def test_stale_or_foreign_scenarios_do_not_produce_proposals(admin_client):
    scope = _scope(admin_client)
    scenario = admin_client.post(f"/api/v1/quotes/{scope['quoteId']}/scenarios", {"name": "Base"}, format="json").data
    hospital = _benefit(admin_client, "HOSPITALIZATION")
    admin_client.put(
        f"/api/v1/plans/{scope['planId']}/benefits/{hospital['id']}",
        {"configuration": {"limits": [{"kind": "ANNUAL_PER_PERSON", "amount": "20000000"}]}}, format="json",
    )
    stale = _generate(admin_client, scope, scenarioId=scenario["id"])
    assert stale.status_code == 409 and "recalculate" in stale.data["detail"]
    admin_client.post(f"/api/v1/scenarios/{scenario['id']}/recalculate", {}, format="json")
    assert _generate(admin_client, scope, scenarioId=scenario["id"]).status_code == 201
    import uuid

    assert _generate(admin_client, scope, scenarioId=str(uuid.uuid4())).status_code == 400
    assert _generate(admin_client, scope, format="XLS").status_code == 400


@pytest.mark.django_db
def test_custom_template_controls_title_sections_and_texts(admin_client):
    scope = _scope(admin_client)
    body = {"title": "Proposta {{quote.number}} para {{client.name}}", "intro": "Prémio total: {{premium.total}} ({{insured.count}} pessoas).",
            "footer": "Rodapé {{document.number}}", "sections": ["premium", "summary"]}
    created = admin_client.post(
        "/api/v1/document-templates",
        {"templateCode": "QUOTE_PROPOSAL_DEFAULT", "storageRef": json.dumps(body), "effectiveFrom": "2026-01-01"}, format="json",
    ).data
    preview = admin_client.get(f"/api/v1/document-templates/{created['id']}/preview")
    assert preview.status_code == 200 and "Cliente de Exemplo" in preview.content.decode("utf-8")
    admin_client.post(f"/api/v1/document-templates/{created['id']}/publish")

    html = _html(admin_client, scope)
    assert "Proposta " in html and "para Holder" in html and "Prémio total: 200 Kz (2 pessoas)." in html
    assert "Rodapé (a atribuir)" in html
    assert html.index("Prémio") < html.index("Resumo da proposta")  # secções pela ordem do GTE
    assert "Coberturas, capitais" not in html and "Exclusões" not in html
    doc = _generate(admin_client, scope, format="HTML").data
    final = admin_client.get(doc["downloadUrl"]).content.decode("utf-8")
    assert f"Rodapé {doc['documentNumber']}" in final and doc["templateVersion"] == 1


@pytest.mark.django_db
def test_invalid_template_bodies_are_rejected_and_legacy_ones_accepted(admin_client):
    def create(body):
        return admin_client.post(
            "/api/v1/document-templates",
            {"templateCode": "T", "storageRef": body, "effectiveFrom": "2026-01-01"}, format="json",
        )

    assert create(json.dumps({"title": "x {{nope.var}}"})).status_code == 400
    assert create(json.dumps({"sections": ["summary", "magic"]})).status_code == 400
    assert create(json.dumps({"sections": []})).status_code == 400
    assert create(json.dumps({"sections": ["summary", "summary"]})).status_code == 400
    assert create(json.dumps({"colour": "red"})).status_code == 400
    assert create("{not json").status_code == 400
    assert create("s3://bucket/legacy.docx").status_code == 201  # apontador antigo: desenho por omissão
    schema = admin_client.get("/api/v1/document-templates/schema").data
    assert "premium.total" in schema["variables"] and "coverages" in schema["sections"] and schema["formats"] == ["HTML", "PDF", "DOCX"]
    assert schema["defaultBody"]["sections"][0] == "summary"


@pytest.mark.django_db
def test_cancelled_quotes_have_no_proposals(admin_client):
    scope = _scope(admin_client)
    quote = _quote(admin_client, scope)
    admin_client.post(f"/api/v1/quotes/{scope['quoteId']}/cancel", {"reasonCode": "T", "reasonText": "t", "rowVersion": quote["rowVersion"]}, format="json")
    assert _generate(admin_client, scope).status_code == 409


@pytest.mark.django_db
def test_failed_generation_is_recorded_and_can_be_retried(admin_client, monkeypatch):
    from apps.quote.application import proposal_layout

    scope = _scope(admin_client)
    original = proposal_layout.RENDERERS["PDF"]

    def boom(blocks, *, title):
        raise RuntimeError("falha de renderização simulada")

    monkeypatch.setitem(proposal_layout.RENDERERS, "PDF", (boom, original[1], original[2]))
    failed = _generate(admin_client, scope)
    assert failed.status_code == 200 and failed.data["status"] == "FAILED" and "simulada" in failed.data["failure"]
    assert failed.data["downloadUrl"] is None
    assert admin_client.get(f"/api/v1/documents/{failed.data['id']}/download").status_code == 409

    monkeypatch.setitem(proposal_layout.RENDERERS, "PDF", original)
    retried = admin_client.post(f"/api/v1/documents/{failed.data['id']}/retry", {}, format="json")
    assert retried.status_code == 200 and retried.data["status"] == "GENERATED" and retried.data["documentNumber"]
    assert admin_client.get(retried.data["downloadUrl"]).content.startswith(b"%PDF")
    assert admin_client.post(f"/api/v1/documents/{failed.data['id']}/retry", {}, format="json").status_code == 409


@pytest.mark.django_db
def test_download_detects_a_tampered_file(admin_client):
    scope = _scope(admin_client)
    doc = _generate(admin_client, scope, format="HTML").data
    from apps.quote.infrastructure.models import Document

    path = Path(Document._base_manager.get(id=doc["id"]).storage_ref)
    path.write_bytes(path.read_bytes() + b"<!-- adulterado -->")
    resp = admin_client.get(doc["downloadUrl"])
    assert resp.status_code == 409 and "integrity" in resp.data["detail"]


@pytest.mark.django_db
def test_readers_can_download_but_not_generate(admin_client, tenant):
    scope = _scope(admin_client)
    doc = _generate(admin_client, scope, format="HTML").data
    reader = make_client(tenant_id=tenant.id, roles=[], scope="q360.read q360.audit")
    assert reader.get(f"/api/v1/quotes/{scope['quoteId']}/proposals").status_code == 200
    assert reader.get(doc["downloadUrl"]).status_code == 200
    assert reader.post(f"/api/v1/quotes/{scope['quoteId']}/proposals", {"format": "PDF"}, format="json").status_code == 403
    assert reader.get("/api/v1/document-templates/schema").status_code == 200


@pytest.mark.django_db
def test_proposal_generation_is_audited(admin_client, tenant):
    scope = _scope(admin_client)
    doc = _generate(admin_client, scope).data
    auditor = make_client(tenant_id=tenant.id, roles=[], scope="q360.read q360.audit")
    events = auditor.get(f"/api/v1/audit-events?entityId={doc['id']}").data["items"]
    generated = next(e for e in events if e["action"] == "document.generated")
    assert generated["metadata"]["after"]["format"] == "PDF" and generated["metadata"]["checksum"] == doc["checksum"]
