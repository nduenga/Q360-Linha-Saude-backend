"""Biblioteca de coberturas (R6) segundo o PDF "Coberturas e Capitais".

- Só Hospitalização, Cirurgia de ambulatório e Urgências/ambulância são BASE
  (B em todas as colunas da matriz do PDF): estão em todos os planos e não se
  removem nem se desativam.
- O resto (incluindo o Ambulatório, "O" no Individual Base) é COMPLEMENTARY.
- O GTE cria/edita/desativa coberturas COMPLEMENTARY, com histórico,
  rollback, pesquisa e importação/exportação.
"""
from __future__ import annotations

import io
import json

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile

from apps.product.domain.market_catalog import BASE_CODES, CATALOG, CATALOG_BY_CODE


def _plan(client, *, product_code: str) -> dict:
    product = client.post("/api/v1/products", {"code": product_code, "name": product_code}).data
    plan = client.post("/api/v1/plans", {"productId": product["id"], "code": "BASIC", "name": "Basic"}).data
    return {"productId": product["id"], "planId": plan["id"]}


def _benefit(client, code: str) -> dict:
    items = client.get(f"/api/v1/benefits?q={code}&status=ALL&size=200").data["items"]
    return next(item for item in items if item["code"] == code)


# ------------------------------------------------------------ catálogo (PDF)
@pytest.mark.django_db
def test_benefit_catalog_matches_the_pdf_coverage_list(admin_client):
    resp = admin_client.get("/api/v1/benefits?size=200")
    assert resp.status_code == 200
    assert {item["code"] for item in resp.data["items"]} == {e["code"] for e in CATALOG}
    assert resp.data["page"]["totalElements"] >= 24


@pytest.mark.django_db
def test_pdf_coverages_are_all_present(admin_client):
    """Every coverage named in the PDF's base/optional lists and matrix."""
    codes = {item["code"] for item in admin_client.get("/api/v1/benefits?size=200").data["items"]}
    expected = {
        "HOSPITALIZATION", "AMBULATORY_SURGERY", "EMERGENCY", "OUTPATIENT", "PHARMACY", "MATERNITY", "DENTAL",
        "OPTICAL", "MEDICAL_EVACUATION", "CRITICAL_ILLNESS", "DIALYSIS_ONCOLOGY", "ENDEMIC_DISEASES",
        "PHYSIOTHERAPY", "PSYCHOLOGY_PSYCHIATRY", "CHECKUP_PREVENTIVE", "VACCINATION", "ALTERNATIVE_MEDICINE",
        "PROSTHESIS", "HOSPITAL_DAILY_ALLOWANCE", "ASSISTANCE_24H", "SECOND_OPINION", "OCCUPATIONAL_MEDICINE",
        "PRE_EXISTING_CONDITIONS",
    }
    assert expected <= codes


@pytest.mark.django_db
def test_only_the_three_b_in_every_column_coverages_are_base(admin_client):
    resp = admin_client.get("/api/v1/benefits?category=BASE")
    assert resp.status_code == 200
    assert {item["code"] for item in resp.data["items"]} == {"HOSPITALIZATION", "AMBULATORY_SURGERY", "EMERGENCY"}
    assert BASE_CODES == {"HOSPITALIZATION", "AMBULATORY_SURGERY", "EMERGENCY"}
    # Ambulatório é "O" no Individual Base -> complementar.
    assert _benefit(admin_client, "OUTPATIENT")["category"] == "COMPLEMENTARY"


@pytest.mark.django_db
def test_catalog_definition_carries_the_pdf_rules(admin_client):
    maternity = _benefit(admin_client, "MATERNITY")["definition"]
    assert maternity["guaranteeForm"] == "PER_EVENT"
    assert maternity["typicalWaitingDays"] == {"individual": [300, 365], "group": [0, 270]}
    assert maternity["antiSelectionWaiting"] is True

    evacuation = _benefit(admin_client, "MEDICAL_EVACUATION")["definition"]
    assert evacuation["requiresPriorApproval"] is True
    assert evacuation["allowedCurrencies"] == ["AOA", "USD"]
    assert evacuation["destinations"] == ["Portugal", "África do Sul", "Brasil", "Namíbia"]

    critical = _benefit(admin_client, "CRITICAL_ILLNESS")["definition"]
    assert critical["guaranteeForm"] == "FIXED_AMOUNT"
    assert "Transplante" in critical["closedDiseaseList"]

    allowance = _benefit(admin_client, "HOSPITAL_DAILY_ALLOWANCE")["definition"]
    assert allowance["typicalMaxDaysPerYear"] == [30, 60]
    assert allowance["timeDeductibleDays"] == [2, 3]

    assistance = _benefit(admin_client, "ASSISTANCE_24H")["definition"]
    assert assistance["noParticipation"] is True
    assert all(e["definition"]["indicative"] for e in CATALOG)


# ----------------------------------------------------- associação a planos
@pytest.mark.django_db
def test_new_plan_automatically_gets_the_3_base_coverages(admin_client):
    refs = _plan(admin_client, product_code="BEN-PROD-1")
    resp = admin_client.get(f"/api/v1/plans/{refs['planId']}/benefits")
    assert resp.status_code == 200
    base_codes = {item["benefit"]["code"] for item in resp.data if item["benefit"]["category"] == "BASE"}
    assert base_codes == {"HOSPITALIZATION", "AMBULATORY_SURGERY", "EMERGENCY"}


@pytest.mark.django_db
def test_complementary_coverage_can_be_added_and_removed(admin_client):
    refs = _plan(admin_client, product_code="BEN-PROD-2")
    pharmacy = _benefit(admin_client, "PHARMACY")

    added = admin_client.post(
        f"/api/v1/plans/{refs['planId']}/benefits", {"benefitId": pharmacy["id"]}, format="json"
    )
    assert added.status_code == 201
    assert added.data["benefit"]["code"] == "PHARMACY"

    listed = admin_client.get(f"/api/v1/plans/{refs['planId']}/benefits").data
    assert any(item["benefit"]["code"] == "PHARMACY" for item in listed)

    removed = admin_client.delete(f"/api/v1/plans/{refs['planId']}/benefits/{pharmacy['id']}")
    assert removed.status_code == 204
    listed_after = admin_client.get(f"/api/v1/plans/{refs['planId']}/benefits").data
    assert not any(item["benefit"]["code"] == "PHARMACY" for item in listed_after)


@pytest.mark.django_db
def test_cannot_remove_a_base_coverage(admin_client):
    refs = _plan(admin_client, product_code="BEN-PROD-3")
    hospitalization = _benefit(admin_client, "HOSPITALIZATION")

    resp = admin_client.delete(f"/api/v1/plans/{refs['planId']}/benefits/{hospitalization['id']}")
    assert resp.status_code == 400
    assert resp.data["code"] == "VALIDATION-400"

    listed = admin_client.get(f"/api/v1/plans/{refs['planId']}/benefits").data
    assert any(item["benefit"]["code"] == "HOSPITALIZATION" for item in listed)


# ------------------------------------------------------------------- CRUD
@pytest.mark.django_db
def test_create_complementary_benefit_and_search(admin_client):
    created = admin_client.post(
        "/api/v1/benefits",
        {
            "code": "TELEDERMATOLOGY",
            "name": "Teledermatologia",
            "description": "Consulta de dermatologia à distância",
            "definition": {"guaranteeForm": "QUANTITY", "typicalSublimits": ["4 consultas/ano"]},
        },
        format="json",
    )
    assert created.status_code == 201
    assert created.data["category"] == "COMPLEMENTARY"
    assert created.data["status"] == "ACTIVE"

    found = admin_client.get("/api/v1/benefits?q=derma").data["items"]
    assert [item["code"] for item in found] == ["TELEDERMATOLOGY"]
    assert admin_client.get(f"/api/v1/benefits/{created.data['id']}").data["name"] == "Teledermatologia"


@pytest.mark.django_db
def test_create_benefit_rejects_bad_code_duplicate_and_bad_definition(admin_client):
    bad_code = admin_client.post("/api/v1/benefits", {"code": "9 bad!", "name": "X"}, format="json")
    assert bad_code.status_code == 400
    # Códigos amigáveis são normalizados para UPPER_SNAKE_CASE.
    friendly = admin_client.post("/api/v1/benefits", {"code": "lower case", "name": "X"}, format="json")
    assert friendly.status_code == 201 and friendly.data["code"] == "LOWER_CASE"

    duplicate = admin_client.post("/api/v1/benefits", {"code": "PHARMACY", "name": "X"}, format="json")
    assert duplicate.status_code == 409

    bad_form = admin_client.post(
        "/api/v1/benefits",
        {"code": "NEW_ONE", "name": "X", "definition": {"guaranteeForm": "NOT_A_FORM"}},
        format="json",
    )
    assert bad_form.status_code == 400

    bad_waiting = admin_client.post(
        "/api/v1/benefits",
        {"code": "NEW_TWO", "name": "X", "definition": {"typicalWaitingDays": {"individual": [90, 30]}}},
        format="json",
    )
    assert bad_waiting.status_code == 400

    bad_currency = admin_client.post(
        "/api/v1/benefits",
        {"code": "NEW_THREE", "name": "X", "definition": {"allowedCurrencies": ["EUR"]}},
        format="json",
    )
    assert bad_currency.status_code == 400


@pytest.mark.django_db
def test_update_benefit_uses_optimistic_locking_and_records_revisions(admin_client):
    dental = _benefit(admin_client, "DENTAL")
    updated = admin_client.patch(
        f"/api/v1/benefits/{dental['id']}",
        {"name": "Estomatologia (revisto)", "description": dental["description"], "definition": dental["definition"],
         "rowVersion": dental["rowVersion"], "reason": "ajuste de nome"},
        format="json",
    )
    assert updated.status_code == 200
    assert updated.data["name"] == "Estomatologia (revisto)"
    assert updated.data["rowVersion"] == dental["rowVersion"] + 1

    stale = admin_client.patch(
        f"/api/v1/benefits/{dental['id']}",
        {"name": "Outra", "rowVersion": dental["rowVersion"]},
        format="json",
    )
    assert stale.status_code == 409
    assert stale.data["code"] == "VERSION-CONFLICT"

    revisions = admin_client.get(f"/api/v1/benefits/{dental['id']}/revisions").data
    assert revisions[0]["action"] == "UPDATED"
    assert revisions[0]["snapshot"]["name"] == "Estomatologia (revisto)"
    assert revisions[0]["reason"] == "ajuste de nome"


@pytest.mark.django_db
def test_deactivate_requires_reason_and_hides_from_default_listing(admin_client):
    vaccination = _benefit(admin_client, "VACCINATION")
    no_reason = admin_client.post(f"/api/v1/benefits/{vaccination['id']}/deactivate", {"reason": "curto"}, format="json")
    assert no_reason.status_code == 400

    done = admin_client.post(
        f"/api/v1/benefits/{vaccination['id']}/deactivate", {"reason": "Já não é comercializada"}, format="json"
    )
    assert done.status_code == 200
    assert done.data["status"] == "INACTIVE"

    active_codes = {i["code"] for i in admin_client.get("/api/v1/benefits?size=200").data["items"]}
    assert "VACCINATION" not in active_codes
    assert _benefit(admin_client, "VACCINATION")["status"] == "INACTIVE"

    back = admin_client.post(
        f"/api/v1/benefits/{vaccination['id']}/reactivate", {"reason": "Voltou a ser comercializada"}, format="json"
    )
    assert back.data["status"] == "ACTIVE"


@pytest.mark.django_db
def test_base_coverage_cannot_be_deactivated(admin_client):
    hospitalization = _benefit(admin_client, "HOSPITALIZATION")
    resp = admin_client.post(
        f"/api/v1/benefits/{hospitalization['id']}/deactivate", {"reason": "tentativa inválida"}, format="json"
    )
    assert resp.status_code == 400
    assert _benefit(admin_client, "HOSPITALIZATION")["status"] == "ACTIVE"


@pytest.mark.django_db
def test_deactivating_a_coverage_in_use_keeps_it_on_existing_plans(admin_client):
    refs = _plan(admin_client, product_code="BEN-PROD-4")
    pharmacy = _benefit(admin_client, "PHARMACY")
    admin_client.post(f"/api/v1/plans/{refs['planId']}/benefits", {"benefitId": pharmacy["id"]}, format="json")

    usage = admin_client.get(f"/api/v1/benefits/{pharmacy['id']}/usage").data
    assert usage["plansUsingBenefit"] == 1
    assert usage["plans"][0]["planCode"] == "BASIC"

    admin_client.post(f"/api/v1/benefits/{pharmacy['id']}/deactivate", {"reason": "descontinuada em novos planos"}, format="json")
    listed = admin_client.get(f"/api/v1/plans/{refs['planId']}/benefits").data
    assert any(item["benefit"]["code"] == "PHARMACY" for item in listed)


@pytest.mark.django_db
def test_rollback_restores_a_previous_revision_as_a_new_revision(admin_client):
    created = admin_client.post(
        "/api/v1/benefits", {"code": "ROLLBACK_ME", "name": "Original", "definition": {}}, format="json"
    ).data
    admin_client.patch(
        f"/api/v1/benefits/{created['id']}",
        {"name": "Alterada", "definition": {}, "rowVersion": created["rowVersion"]},
        format="json",
    )
    revisions = admin_client.get(f"/api/v1/benefits/{created['id']}/revisions").data
    assert [r["revisionNo"] for r in revisions] == [2, 1]

    short = admin_client.post(f"/api/v1/benefits/{created['id']}/rollback", {"revisionNo": 1, "reason": "x"}, format="json")
    assert short.status_code == 400
    missing = admin_client.post(
        f"/api/v1/benefits/{created['id']}/rollback", {"revisionNo": 99, "reason": "revisão inexistente"}, format="json"
    )
    assert missing.status_code == 404

    restored = admin_client.post(
        f"/api/v1/benefits/{created['id']}/rollback", {"revisionNo": 1, "reason": "voltar ao nome original"}, format="json"
    )
    assert restored.status_code == 200
    assert restored.data["name"] == "Original"
    after = admin_client.get(f"/api/v1/benefits/{created['id']}/revisions").data
    assert [r["revisionNo"] for r in after] == [3, 2, 1]
    assert after[0]["action"] == "ROLLBACK_TO_1"


@pytest.mark.django_db
def test_changes_are_audited_with_before_after_and_reason(admin_client, tenant):
    from tests.product.conftest import make_client

    auditor = make_client(tenant_id=tenant.id, roles=["ADMIN"], scope="q360.read q360.audit")
    vaccination = _benefit(admin_client, "VACCINATION")
    admin_client.post(f"/api/v1/benefits/{vaccination['id']}/deactivate", {"reason": "Descontinuada pelo GTE"}, format="json")
    events = auditor.get(f"/api/v1/audit-events?entityType=Benefit&entityId={vaccination['id']}&action=benefit.deactivated")
    assert events.status_code == 200
    metadata = events.data["items"][0]["metadata"]
    assert metadata["reason"] == "Descontinuada pelo GTE"
    assert metadata["before"]["status"] == "ACTIVE"
    assert metadata["after"]["status"] == "INACTIVE"


# --------------------------------------------------------- importar/exportar
@pytest.mark.django_db
def test_export_csv_and_xlsx_roundtrip(admin_client):
    csv_resp = admin_client.get("/api/v1/benefits/export?fileType=csv")
    assert csv_resp.status_code == 200
    text = csv_resp.content.decode("utf-8-sig")
    assert text.splitlines()[0] == "code,name,category,description,status,definition_json"
    assert "HOSPITALIZATION" in text

    xlsx_resp = admin_client.get("/api/v1/benefits/export?fileType=xlsx")
    assert xlsx_resp.status_code == 200
    from openpyxl import load_workbook

    sheet = load_workbook(io.BytesIO(xlsx_resp.content)).active
    assert [c.value for c in sheet[1]][:3] == ["code", "name", "category"]
    assert sheet.max_row >= 25

    assert admin_client.get("/api/v1/benefits/export?fileType=pdf").status_code == 400


@pytest.mark.django_db
def test_import_creates_updates_and_reports_row_errors(admin_client):
    content = (
        "code,name,category,description,status,definition_json\n"
        'NEW_COVER,Nova cobertura,COMPLEMENTARY,Desc,ACTIVE,"{""guaranteeForm"": ""SUBLIMIT""}"\n'
        "PHARMACY,Medicamentos (importado),COMPLEMENTARY,,ACTIVE,\n"
        "9BAD CODE,Erro,COMPLEMENTARY,,ACTIVE,\n"
        "NEW_BASE,Base indevida,BASE,,ACTIVE,\n"
        "BAD_JSON,Json mau,COMPLEMENTARY,,ACTIVE,{not json\n"
        ",Sem codigo,COMPLEMENTARY,,ACTIVE,\n"
    ).encode("utf-8")

    dry = admin_client.post(
        "/api/v1/benefits/import",
        {"file": SimpleUploadedFile("c.csv", content, content_type="text/csv"), "dryRun": "true"},
        format="multipart",
    )
    assert dry.status_code == 200
    assert dry.data["created"] == ["NEW_COVER"]
    assert dry.data["updated"] == ["PHARMACY"]
    assert len(dry.data["errors"]) == 4
    assert admin_client.get("/api/v1/benefits?q=NEW_COVER").data["items"] == []  # nada gravado em dry-run

    real = admin_client.post(
        "/api/v1/benefits/import",
        {"file": SimpleUploadedFile("c.csv", content, content_type="text/csv")},
        format="multipart",
    )
    assert real.status_code == 200
    assert _benefit(admin_client, "NEW_COVER")["definition"] == {"guaranteeForm": "SUBLIMIT"}
    assert _benefit(admin_client, "PHARMACY")["name"] == "Medicamentos (importado)"
    assert {e["row"] for e in real.data["errors"]} == {4, 5, 6, 7}


@pytest.mark.django_db
def test_import_rejects_unknown_file_type_and_missing_file(admin_client):
    assert admin_client.post("/api/v1/benefits/import", {}, format="multipart").status_code == 400
    bad = admin_client.post(
        "/api/v1/benefits/import",
        {"file": SimpleUploadedFile("c.txt", b"x", content_type="text/plain")},
        format="multipart",
    )
    assert bad.status_code == 400


@pytest.mark.django_db
def test_read_only_role_cannot_manage_the_library(tenant):
    from tests.product.conftest import make_client

    reader = make_client(tenant_id=tenant.id, roles=[], scope="q360.read q360.audit")
    assert reader.get("/api/v1/benefits").status_code == 200
    assert reader.post("/api/v1/benefits", {"code": "NOPE_X", "name": "x"}, format="json").status_code == 403
