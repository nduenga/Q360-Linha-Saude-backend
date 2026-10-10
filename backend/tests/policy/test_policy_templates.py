"""A apólice digital usa o modelo `POLICY_DEFAULT` publicado: papel timbrado (PDF) e
modelo Word com marcas, sem JSON."""
from __future__ import annotations

import datetime
import io

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile

from tests.policy.test_issuance import _approved_quote, _approver
from tests.quote.test_template_files import DOCX, _letterhead_pdf, _text_of_pdf, _word

_SIGN = {
    "signatoryName": "Holder", "signatoryDocumentType": "BI", "signatoryDocumentNumber": "000000000BI000",
    "signatoryCapacity": "TOMADOR", "termMonths": 12, "paymentFrequency": "MENSAL", "signedAt": "2026-06-10T14:00:00Z",
}
_PAYMENT = {"paidAt": "2026-06-10T15:00:00Z", "paymentMethod": "TRANSFERENCIA", "paymentReference": "TRF-77"}


def _publish_policy_template(client):
    template = client.post(
        "/api/v1/document-templates",
        {"documentType": "POLICY", "effectiveFrom": datetime.date.today().isoformat()}, format="json",
    ).data
    base = f"/api/v1/document-templates/{template['id']}"
    assert client.post(
        f"{base}/letterhead",
        {"file": SimpleUploadedFile("t.pdf", _letterhead_pdf("TIMBRADO APOLICE"), content_type="application/pdf")}, format="multipart",
    ).status_code == 200
    word = _word(
        paragraphs=["APÓLICE {{policy.number}} de {{client.name}} - {{policy.startDate}} a {{policy.endDate}}", "Pago por {{payment.firstMethod}} ({{payment.firstReference}})"],
        table_row=["{{installment.no}}", "{{installment.dueDate}}"],
    )
    uploaded = client.post(f"{base}/word", {"file": SimpleUploadedFile("a.docx", word, content_type=DOCX)}, format="multipart")
    assert uploaded.status_code == 200, uploaded.data
    assert client.post(f"{base}/publish").status_code == 200
    return template


@pytest.mark.django_db
def test_policy_pdf_uses_the_letterhead_and_docx_uses_the_word_template(admin_client, tenant):
    _publish_policy_template(admin_client)
    quote = _approved_quote(admin_client, tenant, suffix="T1")
    policy = _approver(tenant).post(
        f"/api/v1/quotes/{quote['id']}/policy/emit", {**_SIGN, "rowVersion": quote["rowVersion"], "payment": _PAYMENT}, format="json"
    ).data
    assert policy["status"] == "ACTIVE"

    pdf = admin_client.get(policy["documentPath"])
    assert pdf.status_code == 200 and pdf["Content-Type"] == "application/pdf"
    first_page = _text_of_pdf(pdf.content)[0]
    assert "TIMBRADO APOLICE" in first_page and policy["policyNumber"] in first_page

    word = admin_client.get(policy["documentPath"], {"fileType": "docx"})
    assert word.status_code == 200 and word["Content-Type"] == DOCX
    from docx import Document

    document = Document(io.BytesIO(word.content))
    texts = [p.text for p in document.paragraphs]
    assert any(t.startswith(f"APÓLICE {policy['policyNumber']} de Holder - 2026-06-11 a 2027-06-10") for t in texts)
    assert any("Transferência bancária (TRF-77)" in t for t in texts)
    assert len(document.tables[0].rows) == 13  # cabeçalho + 12 prestações mensais
    assert all("{{" not in cell.text for row in document.tables[0].rows for cell in row.cells)


@pytest.mark.django_db
def test_policy_without_a_template_keeps_the_default_layout(admin_client, tenant):
    quote = _approved_quote(admin_client, tenant, suffix="T2")
    policy = _approver(tenant).post(
        f"/api/v1/quotes/{quote['id']}/policy/emit", {**_SIGN, "rowVersion": quote["rowVersion"], "payment": _PAYMENT}, format="json"
    ).data
    first_page = _text_of_pdf(admin_client.get(policy["documentPath"]).content)[0]
    assert policy["policyNumber"] in first_page and "TIMBRADO" not in first_page
    docx = admin_client.get(policy["documentPath"], {"fileType": "docx"})
    assert docx.status_code == 200 and docx["Content-Type"] == DOCX  # Word gerado pelo sistema
